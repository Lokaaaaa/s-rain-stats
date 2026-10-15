import os, json, sqlite3, time, threading, re, hashlib
from datetime import datetime, timezone
from collections import defaultdict
from typing import List, Optional

from fastapi import FastAPI, HTTPException, Header, UploadFile, File, Form
from fastapi.responses import HTMLResponse, FileResponse, JSONResponse
from pydantic import BaseModel

import requests
from bs4 import BeautifulSoup

try:
    from docx import Document
    from docx.shared import Pt
    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False

DB_PATH = os.environ.get("DB_PATH", "shuffle.db")
DATABASE_URL = os.environ.get("DATABASE_URL", "")
USE_POSTGRES = bool(DATABASE_URL)
API_KEY = os.environ.get("API_KEY", "CHANGE_ME_SECRET_KEY")
GROQ_API_KEY = os.environ.get("GROQ_API_KEY", "")
GROQ_MODEL = "openai/gpt-oss-20b"
MY_NICK = os.environ.get("MY_NICK", "Merami")

SERVER_START_TS = time.time()
SERVER_GIT_COMMIT = (os.environ.get("RENDER_GIT_COMMIT") or "")[:7] or "?"
SERVER_GIT_BRANCH = os.environ.get("RENDER_GIT_BRANCH") or "-"
SERVER_SERVICE_NAME = os.environ.get("RENDER_SERVICE_NAME") or "?"

_default_screenshots = os.path.join(os.path.dirname(os.path.abspath(__file__)), "screenshots")
SCREENSHOT_DIR = os.environ.get("SCREENSHOT_DIR", _default_screenshots)
try:
    os.makedirs(SCREENSHOT_DIR, exist_ok=True)
except Exception:
    SCREENSHOT_DIR = _default_screenshots
    os.makedirs(SCREENSHOT_DIR, exist_ok=True)

if USE_POSTGRES:
    import psycopg2
    import psycopg2.extras

app = FastAPI(title="Shuffle Rain & Tips Stats")
_state_lock = threading.Lock()
INSTANCE_STATES = {}
SCREENSHOT_REQUESTS = {}
PROMO_CHANNEL = os.environ.get("PROMO_CHANNEL", "anyclaimershuffle")
PROMO_PARSE_INTERVAL = 5 * 60
BET_TO_RAIN_WINDOW_SEC = 900

LANG_MAP = {
    'ENGLISH': 'английский', 'RUSSIAN': 'русский', 'JAPANESE': 'японский',
    'TURKISH': 'турецкий', 'FRENCH': 'французский', 'SPANISH': 'испанский',
    'PORTUGUESE': 'португальский', 'KOREAN': 'корейский', 'VIETNAMESE': 'вьетнамский',
    'POLISH': 'польский', 'INDONESIAN': 'индонезийский', 'CHINESE': 'китайский',
}
ALL_CHATS = list(LANG_MAP.keys())
CURRENCIES = ["USDT","USDC","BTC","ETH","SOL","LTC","DOGE","XRP","TRX","BNB","SHIB","SHFL","POL","MATIC",
              "AVAX","TON","GRAM","BONK","WIF","PUMP","TRUMP","DAI","CASH","USD"]

# ==================== RATES ====================
RATES_CACHE_FILE = os.path.join(os.path.dirname(os.path.abspath(__file__)), "server_rates.json")
RATE_ALIASES = {"TON":"GRAM","MATIC":"POL","POLYGON":"POL","USD":"USDT","USDC":"USDT","DAI":"USDT","BUSD":"USDT","TUSD":"USDT"}
_server_rates = {}
_server_rates_lock = threading.Lock()
_server_rates_updated = [0.0]


def _load_server_rates():
    global _server_rates
    if os.path.exists(RATES_CACHE_FILE):
        try:
            with open(RATES_CACHE_FILE, "r", encoding="utf-8") as f:
                data = json.load(f)
            with _server_rates_lock:
                _server_rates = data.get("rates", {})
                _server_rates_updated[0] = data.get("last_update", 0)
        except Exception as e:
            print(f"[RATES] {e}")


def _save_server_rates():
    try:
        with _server_rates_lock:
            data = {"rates": _server_rates, "last_update": _server_rates_updated[0]}
        with open(RATES_CACHE_FILE, "w", encoding="utf-8") as f:
            json.dump(data, f, ensure_ascii=False, indent=2)
    except Exception:
        pass


def _fetch_server_rates():
    try:
        url = "https://api.coingecko.com/api/v3/simple/price"
        ids = ("bitcoin,ethereum,solana,litecoin,dogecoin,ripple,tron,binancecoin,tether,usd-coin,shiba-inu,"
               "shuffle-2,polygon-ecosystem-token,avalanche-2,the-open-network,bonk,dogwifcoin,pump-fun,official-trump,dai")
        r = requests.get(url, params={"ids": ids, "vs_currencies": "rub,usd"}, timeout=15)
        if r.status_code != 200:
            return
        data = r.json()
        mapping = {"BTC":"bitcoin","ETH":"ethereum","SOL":"solana","LTC":"litecoin","DOGE":"dogecoin","XRP":"ripple",
                   "TRX":"tron","BNB":"binancecoin","USDT":"tether","USDC":"usd-coin","SHIB":"shiba-inu",
                   "SHFL":"shuffle-2","POL":"polygon-ecosystem-token","AVAX":"avalanche-2","GRAM":"the-open-network",
                   "BONK":"bonk","WIF":"dogwifcoin","PUMP":"pump-fun","TRUMP":"official-trump","DAI":"dai"}
        new = {}
        for sym, cg in mapping.items():
            if cg in data:
                if "rub" in data[cg]:
                    new[sym] = data[cg]["rub"]
                if "usd" in data[cg]:
                    new[f"{sym}_USD"] = data[cg]["usd"]
        if "GRAM" in new:
            new["TON"] = new["GRAM"]
        if "GRAM_USD" in new:
            new["TON_USD"] = new["GRAM_USD"]
        if "POL" in new:
            new["MATIC"] = new["POL"]
        if "POL_USD" in new:
            new["MATIC_USD"] = new["POL_USD"]
        new["CASH"] = new.get("USDT", 95.0)
        new["CASH_USD"] = 1.0
        with _server_rates_lock:
            _server_rates.update(new)
            _server_rates_updated[0] = time.time()
        _save_server_rates()
        print(f"[RATES] Обновлено {len(new)} курсов")
    except Exception as e:
        print(f"[RATES] fetch: {e}")


def server_rate_rub(symbol):
    if not symbol:
        return None
    s = symbol.upper().strip()
    s = RATE_ALIASES.get(s, s)
    with _server_rates_lock:
        return _server_rates.get(s)


def server_rate_usd(symbol):
    if not symbol:
        return None
    s = symbol.upper().strip()
    s = RATE_ALIASES.get(s, s)
    with _server_rates_lock:
        return _server_rates.get(f"{s}_USD")


def rates_refresh_thread():
    _fetch_server_rates()
    while True:
        time.sleep(3600)
        try:
            _fetch_server_rates()
        except Exception:
            pass


# ==================== БАЗА ====================
def get_db():
    if USE_POSTGRES:
        url = DATABASE_URL
        if url.startswith("postgres://"):
            url = url.replace("postgres://", "postgresql://", 1)
        if "sslmode" not in url:
            sep = "&" if "?" in url else "?"
            url += f"{sep}sslmode=require"
        return psycopg2.connect(url)
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    return conn


def ph():
    return "%s" if USE_POSTGRES else "?"


def query_all(sql, params=()):
    sql_pg = sql.replace("?", "%s") if USE_POSTGRES else sql
    conn = get_db()
    try:
        cur = conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor) if USE_POSTGRES else conn.cursor()
        cur.execute(sql_pg, params)
        try:
            return [dict(r) for r in cur.fetchall()]
        except Exception:
            return []
    finally:
        conn.close()


def query_one(sql, params=()):
    sql_pg = sql.replace("?", "%s") if USE_POSTGRES else sql
    conn = get_db()
    try:
        cur = conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor) if USE_POSTGRES else conn.cursor()
        cur.execute(sql_pg, params)
        row = cur.fetchone()
        return dict(row) if row else None
    finally:
        conn.close()


def execute(sql, params=()):
    sql_pg = sql.replace("?", "%s") if USE_POSTGRES else sql
    conn = get_db()
    try:
        cur = conn.cursor()
        cur.execute(sql_pg, params)
        conn.commit()
    finally:
        conn.close()


def _try(sql):
    try:
        execute(sql)
    except Exception:
        pass


def init_db():
    if USE_POSTGRES:
        execute("""CREATE TABLE IF NOT EXISTS events (
            id SERIAL PRIMARY KEY, type TEXT NOT NULL, chat TEXT NOT NULL,
            sender TEXT NOT NULL, receivers TEXT NOT NULL, receivers_info TEXT,
            amount_text TEXT, amount_type TEXT, crypto_symbol TEXT,
            amount_rub DOUBLE PRECISION DEFAULT 0, amount_usd DOUBLE PRECISION DEFAULT 0,
            created_at TEXT NOT NULL, ts DOUBLE PRECISION)""")
    else:
        execute("""CREATE TABLE IF NOT EXISTS events (
            id INTEGER PRIMARY KEY AUTOINCREMENT, type TEXT NOT NULL, chat TEXT NOT NULL,
            sender TEXT NOT NULL, receivers TEXT NOT NULL, receivers_info TEXT,
            amount_text TEXT, amount_type TEXT, crypto_symbol TEXT,
            amount_rub REAL DEFAULT 0, amount_usd REAL DEFAULT 0,
            created_at TEXT NOT NULL, ts REAL)""")
    for col, typ in [
        ("ts", "DOUBLE PRECISION" if USE_POSTGRES else "REAL"),
        ("context_messages", "TEXT"), ("context_bets", "TEXT"),
        ("amount_usd", "DOUBLE PRECISION DEFAULT 0" if USE_POSTGRES else "REAL DEFAULT 0"),
        ("receivers_info", "TEXT"),
    ]:
        _try(f"ALTER TABLE events ADD COLUMN {col} {typ}")
    for idx in ["idx_type ON events(type)", "idx_sender ON events(sender)",
                "idx_events_ts ON events(ts)", "idx_events_chat ON events(chat)"]:
        execute(f"CREATE INDEX IF NOT EXISTS {idx}")

    if USE_POSTGRES:
        execute("""CREATE TABLE IF NOT EXISTS bulk_alerts (
            id SERIAL PRIMARY KEY, ts DOUBLE PRECISION NOT NULL, chat TEXT NOT NULL,
            sender TEXT NOT NULL, amount_text TEXT, amount_rub DOUBLE PRECISION DEFAULT 0,
            crypto TEXT, receivers TEXT NOT NULL, window_sec INTEGER DEFAULT 0,
            is_rubric INTEGER DEFAULT 0)""")
    else:
        execute("""CREATE TABLE IF NOT EXISTS bulk_alerts (
            id INTEGER PRIMARY KEY AUTOINCREMENT, ts REAL NOT NULL, chat TEXT NOT NULL,
            sender TEXT NOT NULL, amount_text TEXT, amount_rub REAL DEFAULT 0,
            crypto TEXT, receivers TEXT NOT NULL, window_sec INTEGER DEFAULT 0,
            is_rubric INTEGER DEFAULT 0)""")
    _try("ALTER TABLE bulk_alerts ADD COLUMN is_rubric INTEGER DEFAULT 0")
    execute("CREATE INDEX IF NOT EXISTS idx_bulk_ts ON bulk_alerts(ts)")
    execute("CREATE INDEX IF NOT EXISTS idx_bulk_sender ON bulk_alerts(sender)")

    if USE_POSTGRES:
        execute("""CREATE TABLE IF NOT EXISTS player_activity (
            id SERIAL PRIMARY KEY, ts DOUBLE PRECISION NOT NULL, username TEXT NOT NULL,
            chat TEXT NOT NULL, text TEXT)""")
    else:
        execute("""CREATE TABLE IF NOT EXISTS player_activity (
            id INTEGER PRIMARY KEY AUTOINCREMENT, ts REAL NOT NULL, username TEXT NOT NULL,
            chat TEXT NOT NULL, text TEXT)""")
    execute("CREATE INDEX IF NOT EXISTS idx_player_ts ON player_activity(ts)")
    execute("CREATE INDEX IF NOT EXISTS idx_player_user ON player_activity(username)")

    if USE_POSTGRES:
        execute("""CREATE TABLE IF NOT EXISTS bets (
            id SERIAL PRIMARY KEY, ts DOUBLE PRECISION NOT NULL, chat TEXT NOT NULL,
            sender TEXT NOT NULL, game TEXT, multiplier TEXT, amount_text TEXT,
            amount_rub DOUBLE PRECISION DEFAULT 0, amount_usd DOUBLE PRECISION DEFAULT 0,
            crypto TEXT, outcome TEXT, source TEXT DEFAULT 'bets_page',
            real_sender TEXT)""")
    else:
        execute("""CREATE TABLE IF NOT EXISTS bets (
            id INTEGER PRIMARY KEY AUTOINCREMENT, ts REAL NOT NULL, chat TEXT NOT NULL,
            sender TEXT NOT NULL, game TEXT, multiplier TEXT, amount_text TEXT,
            amount_rub REAL DEFAULT 0, amount_usd REAL DEFAULT 0, crypto TEXT,
            outcome TEXT, source TEXT DEFAULT 'bets_page', real_sender TEXT)""")
    _try("ALTER TABLE bets ADD COLUMN source TEXT DEFAULT 'bets_page'")
    _try("ALTER TABLE bets ADD COLUMN real_sender TEXT")
    for idx in ["idx_bets_ts ON bets(ts)", "idx_bets_chat ON bets(chat)", "idx_bets_sender ON bets(sender)"]:
        execute(f"CREATE INDEX IF NOT EXISTS {idx}")

    if USE_POSTGRES:
        execute("""CREATE TABLE IF NOT EXISTS intents (
            id SERIAL PRIMARY KEY, ts DOUBLE PRECISION NOT NULL, chat TEXT NOT NULL,
            sender TEXT NOT NULL, text TEXT, percent INTEGER DEFAULT 100,
            matched_rain_id INTEGER, matched_delta_sec INTEGER)""")
    else:
        execute("""CREATE TABLE IF NOT EXISTS intents (
            id INTEGER PRIMARY KEY AUTOINCREMENT, ts REAL NOT NULL, chat TEXT NOT NULL,
            sender TEXT NOT NULL, text TEXT, percent INTEGER DEFAULT 100,
            matched_rain_id INTEGER, matched_delta_sec INTEGER)""")
    _try("ALTER TABLE intents ADD COLUMN percent INTEGER DEFAULT 100")

    if USE_POSTGRES:
        execute("""CREATE TABLE IF NOT EXISTS promocodes (
            id SERIAL PRIMARY KEY, tg_msg_id TEXT UNIQUE, text TEXT NOT NULL, link TEXT,
            message_date TEXT, discovered_at DOUBLE PRECISION, used INTEGER DEFAULT 0,
            used_at DOUBLE PRECISION, used_for_vault INTEGER DEFAULT 0,
            vault_amount_usd DOUBLE PRECISION DEFAULT 0)""")
    else:
        execute("""CREATE TABLE IF NOT EXISTS promocodes (
            id INTEGER PRIMARY KEY AUTOINCREMENT, tg_msg_id TEXT UNIQUE, text TEXT NOT NULL,
            link TEXT, message_date TEXT, discovered_at REAL, used INTEGER DEFAULT 0,
            used_at REAL, used_for_vault INTEGER DEFAULT 0, vault_amount_usd REAL DEFAULT 0)""")
    _try("ALTER TABLE promocodes ADD COLUMN used_for_vault INTEGER DEFAULT 0")
    _try("ALTER TABLE promocodes ADD COLUMN vault_amount_usd DOUBLE PRECISION DEFAULT 0"
         if USE_POSTGRES else
         "ALTER TABLE promocodes ADD COLUMN vault_amount_usd REAL DEFAULT 0")

    if USE_POSTGRES:
        execute("""CREATE TABLE IF NOT EXISTS chat_tasks (
            id SERIAL PRIMARY KEY, chat TEXT NOT NULL, status TEXT DEFAULT 'pending',
            messages TEXT, result TEXT, error TEXT,
            created_at DOUBLE PRECISION, completed_at DOUBLE PRECISION)""")
    else:
        execute("""CREATE TABLE IF NOT EXISTS chat_tasks (
            id INTEGER PRIMARY KEY AUTOINCREMENT, chat TEXT NOT NULL,
            status TEXT DEFAULT 'pending', messages TEXT, result TEXT, error TEXT,
            created_at REAL, completed_at REAL)""")

    if USE_POSTGRES:
        execute("""CREATE TABLE IF NOT EXISTS chat_players (
            id SERIAL PRIMARY KEY, chat TEXT NOT NULL, username TEXT NOT NULL,
            msg_count INTEGER DEFAULT 0, first_seen DOUBLE PRECISION,
            last_seen DOUBLE PRECISION, UNIQUE(chat, username))""")
    else:
        execute("""CREATE TABLE IF NOT EXISTS chat_players (
            id INTEGER PRIMARY KEY AUTOINCREMENT, chat TEXT NOT NULL, username TEXT NOT NULL,
            msg_count INTEGER DEFAULT 0, first_seen REAL, last_seen REAL,
            UNIQUE(chat, username))""")
    for idx in ["idx_cp_chat ON chat_players(chat)", "idx_cp_user ON chat_players(username)",
                "idx_cp_last ON chat_players(last_seen)"]:
        execute(f"CREATE INDEX IF NOT EXISTS {idx}")

    if USE_POSTGRES:
        execute("""CREATE TABLE IF NOT EXISTS player_message_ticks (
            username TEXT NOT NULL, chat TEXT NOT NULL, ts DOUBLE PRECISION NOT NULL,
            PRIMARY KEY (username, chat, ts))""")
    else:
        execute("""CREATE TABLE IF NOT EXISTS player_message_ticks (
            username TEXT NOT NULL, chat TEXT NOT NULL, ts REAL NOT NULL,
            PRIMARY KEY (username, chat, ts))""")
    execute("CREATE INDEX IF NOT EXISTS idx_pmt_user_ts ON player_message_ticks(username, ts)")
    execute("CREATE INDEX IF NOT EXISTS idx_pmt_ts ON player_message_ticks(ts)")

    if USE_POSTGRES:
        execute("""CREATE TABLE IF NOT EXISTS player_profiles (
            id SERIAL PRIMARY KEY, username TEXT UNIQUE NOT NULL,
            lang TEXT DEFAULT 'ENGLISH', description TEXT DEFAULT '',
            is_top BOOLEAN DEFAULT FALSE, updated_at DOUBLE PRECISION)""")
        execute("""CREATE TABLE IF NOT EXISTS account_balances (
            currency TEXT PRIMARY KEY, amount DOUBLE PRECISION DEFAULT 0,
            amount_usd DOUBLE PRECISION DEFAULT 0, updated_at DOUBLE PRECISION)""")
        execute("""CREATE TABLE IF NOT EXISTS vault_balances (
            currency TEXT PRIMARY KEY, amount DOUBLE PRECISION DEFAULT 0,
            amount_usd DOUBLE PRECISION DEFAULT 0, updated_at DOUBLE PRECISION)""")
        execute("""CREATE TABLE IF NOT EXISTS vault_transactions (
            id SERIAL PRIMARY KEY, ts DOUBLE PRECISION NOT NULL, type TEXT NOT NULL,
            currency TEXT, amount DOUBLE PRECISION DEFAULT 0,
            amount_usd DOUBLE PRECISION DEFAULT 0, description TEXT)""")
        execute("""CREATE TABLE IF NOT EXISTS wager_state (
            id INTEGER PRIMARY KEY, amount_required DOUBLE PRECISION DEFAULT 0,
            deadline_ts DOUBLE PRECISION DEFAULT 0, updated_at DOUBLE PRECISION)""")
    else:
        execute("""CREATE TABLE IF NOT EXISTS player_profiles (
            id INTEGER PRIMARY KEY AUTOINCREMENT, username TEXT UNIQUE NOT NULL,
            lang TEXT DEFAULT 'ENGLISH', description TEXT DEFAULT '',
            is_top INTEGER DEFAULT 0, updated_at REAL)""")
        execute("""CREATE TABLE IF NOT EXISTS account_balances (
            currency TEXT PRIMARY KEY, amount REAL DEFAULT 0,
            amount_usd REAL DEFAULT 0, updated_at REAL)""")
        execute("""CREATE TABLE IF NOT EXISTS vault_balances (
            currency TEXT PRIMARY KEY, amount REAL DEFAULT 0,
            amount_usd REAL DEFAULT 0, updated_at REAL)""")
        execute("""CREATE TABLE IF NOT EXISTS vault_transactions (
            id INTEGER PRIMARY KEY AUTOINCREMENT, ts REAL NOT NULL, type TEXT NOT NULL,
            currency TEXT, amount REAL DEFAULT 0, amount_usd REAL DEFAULT 0,
            description TEXT)""")
        execute("""CREATE TABLE IF NOT EXISTS wager_state (
            id INTEGER PRIMARY KEY, amount_required REAL DEFAULT 0,
            deadline_ts REAL DEFAULT 0, updated_at REAL)""")
    execute("CREATE INDEX IF NOT EXISTS idx_vt_ts ON vault_transactions(ts)")
    execute("CREATE INDEX IF NOT EXISTS idx_vt_type ON vault_transactions(type)")
    execute("CREATE TABLE IF NOT EXISTS vault_meta (k TEXT PRIMARY KEY, v TEXT)")
    execute("CREATE TABLE IF NOT EXISTS meta (k TEXT PRIMARY KEY, v TEXT)")

    if USE_POSTGRES:
        execute("""CREATE TABLE IF NOT EXISTS moderators (
            username TEXT PRIMARY KEY, first_seen DOUBLE PRECISION,
            last_seen DOUBLE PRECISION, chats TEXT)""")
        execute("""CREATE TABLE IF NOT EXISTS moderator_messages (
            id SERIAL PRIMARY KEY, ts DOUBLE PRECISION NOT NULL,
            username TEXT NOT NULL, chat TEXT NOT NULL, text TEXT)""")
    else:
        execute("""CREATE TABLE IF NOT EXISTS moderators (
            username TEXT PRIMARY KEY, first_seen REAL,
            last_seen REAL, chats TEXT)""")
        execute("""CREATE TABLE IF NOT EXISTS moderator_messages (
            id INTEGER PRIMARY KEY AUTOINCREMENT, ts REAL NOT NULL,
            username TEXT NOT NULL, chat TEXT NOT NULL, text TEXT)""")
    execute("CREATE INDEX IF NOT EXISTS idx_mm_ts ON moderator_messages(ts)")
    execute("CREATE INDEX IF NOT EXISTS idx_mm_user ON moderator_messages(username)")


def meta_get(k, default=None):
    r = query_one("SELECT v FROM meta WHERE k=?", (k,))
    return r["v"] if r else default


def meta_set(k, v):
    if USE_POSTGRES:
        execute("INSERT INTO meta (k,v) VALUES (%s,%s) ON CONFLICT (k) DO UPDATE SET v=EXCLUDED.v",
                (k, str(v)))
    else:
        execute("INSERT OR REPLACE INTO meta (k,v) VALUES (?,?)", (k, str(v)))


def vault_meta_get(k, default=None):
    r = query_one("SELECT v FROM vault_meta WHERE k=?", (k,))
    return r["v"] if r else default


def vault_meta_set(k, v):
    if USE_POSTGRES:
        execute("INSERT INTO vault_meta (k,v) VALUES (%s,%s) ON CONFLICT (k) DO UPDATE SET v=EXCLUDED.v",
                (k, str(v)))
    else:
        execute("INSERT OR REPLACE INTO vault_meta (k,v) VALUES (?,?)", (k, str(v)))


def migrate_ts_once():
    try:
        if meta_get("ts_migrated_v1"):
            return
        if USE_POSTGRES:
            execute("""UPDATE events SET ts = EXTRACT(EPOCH FROM created_at::timestamptz)
                       WHERE ts IS NULL AND created_at IS NOT NULL""")
        else:
            execute("""UPDATE events SET ts = CAST(strftime('%s', created_at) AS REAL)
                       WHERE ts IS NULL AND created_at IS NOT NULL""")
        meta_set("ts_migrated_v1", "1")
        print("[MIGRATE] ts миграция выполнена")
    except Exception as e:
        print(f"[MIGRATE] {e}")


def dedup_thread():
    while True:
        time.sleep(600)
        try:
            rows = query_all("""SELECT id, amount_text, crypto_symbol, amount_rub, amount_usd
                                FROM events
                                WHERE (amount_rub IS NULL OR amount_rub=0
                                       OR amount_usd IS NULL OR amount_usd=0)
                                LIMIT 5000""")
            for r in rows:
                txt = r.get("amount_text") or ""
                crypto = (r.get("crypto_symbol") or "").upper().strip()
                value = None
                try:
                    value = float(txt.replace("\u00a0", " ").replace(" ", "").replace(",", ".")
                                  .replace("₽", "").replace("$", "").replace("€", ""))
                except Exception:
                    pass
                if value is None:
                    continue
                new_rub = r.get("amount_rub") or 0
                new_usd = r.get("amount_usd") or 0
                has_fiat = any(m in txt for m in ("₽", "$", "€", "£"))
                if crypto and not has_fiat:
                    cr = server_rate_rub(crypto)
                    cu = server_rate_usd(crypto)
                    if cr and not new_rub:
                        new_rub = value * cr
                    if cu and not new_usd:
                        new_usd = value * cu
                elif "₽" in txt and not new_rub:
                    new_rub = value
                elif "$" in txt and not new_usd:
                    new_usd = value
                if new_rub != (r.get("amount_rub") or 0) or new_usd != (r.get("amount_usd") or 0):
                    p = ph()
                    execute(f"UPDATE events SET amount_rub={p}, amount_usd={p} WHERE id={p}",
                            (new_rub, new_usd, r["id"]))
        except Exception:
            pass


def ticks_cleanup_thread():
    while True:
        time.sleep(6 * 3600)
        try:
            execute(f"DELETE FROM player_message_ticks WHERE ts < {ph()}",
                    (time.time() - 7 * 86400,))
        except Exception:
            pass


# ==================== МОДЕЛИ ====================
class EventIn(BaseModel):
    type: str
    chat: str
    sender: str
    receivers: List[str] = []
    receivers_info: List[dict] = []
    amount_text: str = ""
    amount_type: str = "fiat"
    crypto_symbol: str = ""
    amount_rub: float = 0.0
    amount_usd: float = 0.0
    created_at: Optional[str] = None
    ts: Optional[float] = None
    context_messages: List[str] = []
    context_bets: List[dict] = []


class BulkAlertIn(BaseModel):
    ts: Optional[float] = None
    chat: str
    sender: str
    amount_text: str = ""
    amount_rub: float = 0.0
    crypto: str = ""
    receivers: List[str] = []
    window_sec: int = 0


class PlayerActivityIn(BaseModel):
    ts: Optional[float] = None
    username: str
    chat: str
    text: str = ""


class IntentIn(BaseModel):
    ts: Optional[float] = None
    chat: str
    sender: str
    text: str = ""
    percent: int = 100


class BetIn(BaseModel):
    ts: Optional[float] = None
    chat: str
    sender: str
    real_sender: Optional[str] = None
    game: str = ""
    multiplier: str = ""
    amount_text: str = ""
    amount_rub: float = 0.0
    amount_usd: float = 0.0
    crypto: str = ""
    outcome: str = "win"
    source: str = "bets_page"


class InstanceStatus(BaseModel):
    iid: int
    mode: Optional[str] = None
    status: Optional[str] = None
    current_chat: Optional[str] = None
    chats: Optional[List[str]] = None
    new_count: Optional[int] = 0
    threshold: Optional[int] = None
    error: Optional[str] = None
    updated_at: Optional[float] = None


class StatusPayload(BaseModel):
    instances: List[InstanceStatus] = []
    screenshot_requests: List[int] = []


class ChatTaskIn(BaseModel):
    chat: str


class ChatTaskCompleteIn(BaseModel):
    task_id: int
    messages: List[dict] = []
    error: str = ""


class PlayerStatItem(BaseModel):
    name: str
    msgs: int = 0
    last_seen: float = 0.0


class PlayerTick(BaseModel):
    username: str
    ts: float


class PlayersBatchIn(BaseModel):
    chat: str
    players: List[PlayerStatItem] = []
    ticks: List[PlayerTick] = []


class BalanceIn(BaseModel):
    currency: str
    amount: float
    amount_usd: Optional[float] = None


class BalanceDeleteIn(BaseModel):
    currency: str


class VaultTransactionIn(BaseModel):
    type: str
    currency: str = "USDT"
    amount: float = 0.0
    amount_usd: float = 0.0
    description: str = ""
    ts: Optional[float] = None


class VaultWagerIn(BaseModel):
    amount_required: float
    deadline_ts: float
    spent_now: Optional[float] = None


class VaultGoalIn(BaseModel):
    goal_usd: float
    goal_text: str = ""


class TransferIn(BaseModel):
    direction: str
    currency: str
    amount: float


class ControlSet(BaseModel):
    action: str


class NgrokUrlIn(BaseModel):
    url: str
    ts: Optional[float] = None


class ModeratorMessageIn(BaseModel):
    ts: Optional[float] = None
    username: str
    chat: str
    text: str = ""


class LogFilterIn(BaseModel):
    categories: dict


DEFAULT_LOG_FILTER = {
    "rain": True, "tip": True, "bet": False, "bulk": True,
    "intent": True, "player": True, "moderator": True,
    "observer": True, "server": True, "batch": True,
    "rates": False, "shot": False, "proc": True,
    "ngrok": True, "telegram": True, "debug": False,
    "error": True, "info": True,
}


# ==================== РАСЧЁТ ====================
def _recalc_amounts(amount_text, amount_type, crypto_symbol):
    txt = amount_text or ""
    crypto = (crypto_symbol or "").upper().strip()
    cleaned = (txt.replace("\u00a0", " ").replace(" ", "")
               .replace("₽", "").replace("$", "").replace("€", ""))
    if "," in cleaned and "." in cleaned:
        if cleaned.rfind(",") > cleaned.rfind("."):
            cleaned = cleaned.replace(".", "").replace(",", ".")
        else:
            cleaned = cleaned.replace(",", "")
    else:
        cleaned = cleaned.replace(",", ".")
    try:
        value = float(cleaned)
    except Exception:
        return 0.0, 0.0
    has_fiat = any(m in txt for m in ("₽", "$", "€", "£", "₴", "₸", "₺"))
    if crypto and (amount_type == "crypto" or not has_fiat):
        cr = server_rate_rub(crypto) or 0.0
        cu = server_rate_usd(crypto) or 0.0
        return value * cr, value * cu
    if "₽" in txt:
        ur = server_rate_rub("USDT") or 95.0
        uu = server_rate_usd("USDT") or 1.0
        return value, value / ur * uu
    if "$" in txt:
        ur = server_rate_rub("USDT") or 95.0
        uu = server_rate_usd("USDT") or 1.0
        return value / uu * ur, value
    if "€" in txt:
        usd = value * 1.08
        ur = server_rate_rub("USDT") or 95.0
        uu = server_rate_usd("USDT") or 1.0
        return usd / uu * ur, usd
    ur = server_rate_rub("USDT") or 95.0
    uu = server_rate_usd("USDT") or 1.0
    return value, value / ur * uu


def _enrich_receivers_info_from_db(chat_name, receivers_info, event_ts=None):
    """Обогащает receivers_info данными из player_message_ticks.
    Считает 'minutes_ago' ОТНОСИТЕЛЬНО event_ts (времени дождя), а не сейчас.
    Берём последнее сообщение игрока ДО момента дождя (ts <= event_ts).
    """
    if not receivers_info:
        return
    if event_ts is None:
        event_ts = time.time()
    for r in receivers_info:
        name = (r or {}).get("name")
        if not name:
            continue

        # 1. Ищем последнее сообщение игрока ДО момента дождя
        tick_ts = None
        try:
            tick = query_one(
                """SELECT MAX(ts) AS last_before FROM player_message_ticks
                   WHERE username=? AND chat=? AND ts <= ?""",
                (name, chat_name, event_ts)
            )
            if tick and tick.get("last_before"):
                tick_ts = float(tick["last_before"])
        except Exception:
            pass

        # 2. Фолбэк на chat_players.last_seen, если оно МЕНЬШЕ event_ts
        if tick_ts is None:
            try:
                crow = query_one("SELECT last_seen, msg_count FROM chat_players WHERE username=? AND chat=?",
                                 (name, chat_name))
                if crow and crow.get("last_seen"):
                    ls = float(crow["last_seen"])
                    if ls <= event_ts:
                        tick_ts = ls
            except Exception:
                pass

        # 3. Считаем минуты назад ОТ ВРЕМЕНИ ДОЖДЯ
        if tick_ts is not None:
            min_ago = int((event_ts - tick_ts) / 60)
            r["last_msg_ts"] = tick_ts
            r["minutes_ago"] = min_ago
            r["wrote_in_25min"] = min_ago <= 25
            r["in_chat_only"] = False
        else:
            if r.get("minutes_ago") is None:
                r["wrote_in_25min"] = False

        # 4. Обновляем total_msgs если сервер знает больше
        try:
            row = query_one("SELECT msg_count FROM chat_players WHERE username=? AND chat=?",
                            (name, chat_name))
            if row and (row.get("msg_count") or 0) > (r.get("total_msgs") or 0):
                r["total_msgs"] = row["msg_count"]
        except Exception:
            pass


# ==================== API SCRIPT ====================
@app.post("/api/event")
async def add_event(event: EventIn, x_api_key: str = Header(default="")):
    if x_api_key != API_KEY:
        raise HTTPException(401, "Invalid API key")
    if event.type not in ("rain", "tip"):
        raise HTTPException(400, "type must be rain or tip")
    ts_val = event.ts if event.ts else time.time()
    created = event.created_at or datetime.fromtimestamp(ts_val, tz=timezone.utc).isoformat()
    dup = query_one("""SELECT id FROM events
                       WHERE type=? AND chat=? AND sender=? AND receivers=? AND amount_text=?
                       AND ts IS NOT NULL AND ABS(ts-?)<30 LIMIT 1""",
                    (event.type, event.chat, event.sender,
                     json.dumps(event.receivers or [], ensure_ascii=False),
                     event.amount_text, ts_val))
    if dup:
        return {"ok": True, "dup": True}
    if event.type == "rain" and event.receivers_info:
        _enrich_receivers_info_from_db(event.chat, event.receivers_info, event_ts=ts_val)
    rub = event.amount_rub or 0.0
    usd = event.amount_usd or 0.0
    if rub == 0 or usd == 0:
        cr, cu = _recalc_amounts(event.amount_text, event.amount_type, event.crypto_symbol)
        if not rub:
            rub = cr
        if not usd:
            usd = cu
    p = ph()
    execute(f"""INSERT INTO events
                (type,chat,sender,receivers,receivers_info,amount_text,amount_type,
                 crypto_symbol,amount_rub,amount_usd,created_at,ts,
                 context_messages,context_bets)
                VALUES ({p},{p},{p},{p},{p},{p},{p},{p},{p},{p},{p},{p},{p},{p})""",
            (event.type, event.chat, event.sender,
             json.dumps(event.receivers or [], ensure_ascii=False),
             json.dumps(event.receivers_info or [], ensure_ascii=False),
             event.amount_text, event.amount_type, event.crypto_symbol,
             rub, usd, created, ts_val,
             json.dumps(event.context_messages or [], ensure_ascii=False),
             json.dumps(event.context_bets or [], ensure_ascii=False)))
    return {"ok": True}


@app.post("/api/events_batch")
async def events_batch(payload: dict, x_api_key: str = Header(default="")):
    if x_api_key != API_KEY:
        raise HTTPException(401, "Invalid API key")
    events = payload.get("events", [])
    inserted = 0
    for e in events:
        try:
            if e.get("type") == "rain" and e.get("receivers_info"):
                _enrich_receivers_info_from_db(e.get("chat"), e.get("receivers_info"),
                                   event_ts=float(e.get("ts", time.time())))
            rub = float(e.get("amount_rub") or 0)
            usd = float(e.get("amount_usd") or 0)
            if rub == 0 or usd == 0:
                cr, cu = _recalc_amounts(e.get("amount_text", ""),
                                         e.get("amount_type", "fiat"),
                                         e.get("crypto_symbol", ""))
                if not rub:
                    rub = cr
                if not usd:
                    usd = cu
            p = ph()
            execute(f"""INSERT INTO events
                        (type,chat,sender,receivers,receivers_info,amount_text,amount_type,
                         crypto_symbol,amount_rub,amount_usd,created_at,ts,
                         context_messages,context_bets)
                        VALUES ({p},{p},{p},{p},{p},{p},{p},{p},{p},{p},{p},{p},{p},{p})""",
                    (e.get("type"), e.get("chat"), e.get("sender"),
                     json.dumps(e.get("receivers", []), ensure_ascii=False),
                     json.dumps(e.get("receivers_info", []), ensure_ascii=False),
                     e.get("amount_text", ""), e.get("amount_type", "fiat"),
                     e.get("crypto_symbol", ""), rub, usd,
                     datetime.fromtimestamp(e.get("ts", time.time()),
                                            tz=timezone.utc).isoformat(),
                     e.get("ts", time.time()),
                     json.dumps(e.get("context_messages", []), ensure_ascii=False),
                     json.dumps(e.get("context_bets", []), ensure_ascii=False)))
            inserted += 1
        except Exception:
            pass
    return {"ok": True, "inserted": inserted}


@app.post("/api/bulk_alert")
async def add_bulk_alert(body: BulkAlertIn, x_api_key: str = Header(default="")):
    if x_api_key != API_KEY:
        raise HTTPException(401, "Invalid API key")
    ts_val = body.ts if body.ts else time.time()
    rub = body.amount_rub or 0.0
    if rub == 0:
        cr, _ = _recalc_amounts(body.amount_text,
                                "crypto" if body.crypto else "fiat",
                                body.crypto)
        rub = cr
    p = ph()
    execute(f"""INSERT INTO bulk_alerts
                (ts,chat,sender,amount_text,amount_rub,crypto,receivers,window_sec)
                VALUES ({p},{p},{p},{p},{p},{p},{p},{p})""",
            (ts_val, body.chat, body.sender, body.amount_text, rub, body.crypto,
             json.dumps(body.receivers, ensure_ascii=False), body.window_sec))
    return {"ok": True}


@app.post("/api/bulk/rubric/{alert_id}")
async def toggle_bulk_rubric(alert_id: int):
    row = query_one("SELECT id, is_rubric FROM bulk_alerts WHERE id=?", (alert_id,))
    if not row:
        raise HTTPException(404, "not found")
    nv = 0 if row.get("is_rubric") else 1
    execute("UPDATE bulk_alerts SET is_rubric=? WHERE id=?", (nv, alert_id))
    return {"ok": True, "is_rubric": nv}


@app.post("/api/player_activity")
async def add_player_activity(body: PlayerActivityIn, x_api_key: str = Header(default="")):
    if x_api_key != API_KEY:
        raise HTTPException(401, "Invalid API key")
    ts_val = body.ts if body.ts else time.time()
    execute(f"INSERT INTO player_activity (ts,username,chat,text) VALUES ({ph()},{ph()},{ph()},{ph()})",
            (ts_val, body.username, body.chat, body.text))
    return {"ok": True}


@app.post("/api/bet")
async def add_bet(body: BetIn, x_api_key: str = Header(default="")):
    if x_api_key != API_KEY:
        raise HTTPException(401, "Invalid API key")
    ts_val = body.ts if body.ts else time.time()
    rub = body.amount_rub or 0.0
    usd = body.amount_usd or 0.0
    if rub == 0 or usd == 0:
        cr, cu = _recalc_amounts(body.amount_text,
                                 "crypto" if body.crypto else "fiat",
                                 body.crypto)
        if not rub:
            rub = cr
        if not usd:
            usd = cu
    p = ph()
    execute(f"""INSERT INTO bets
                (ts,chat,sender,real_sender,game,multiplier,amount_text,
                 amount_rub,amount_usd,crypto,outcome,source)
                VALUES ({p},{p},{p},{p},{p},{p},{p},{p},{p},{p},{p},{p})""",
            (ts_val, body.chat, body.sender, body.real_sender, body.game, body.multiplier,
             body.amount_text, rub, usd, body.crypto, body.outcome,
             body.source or "bets_page"))
    return {"ok": True}


@app.post("/api/bets/clear_top")
async def clear_bets_top(hard: int = 0, x_api_key: str = Header(default="")):
    """
    Очищает "Топ отправителей ставок".
    По умолчанию (hard=0): ставит timestamp, /api/bets_top_senders игнорирует ставки до него.
    hard=1: реально удаляет все записи из таблицы bets.
    """
    if x_api_key != API_KEY:
        raise HTTPException(401, "Invalid API key")
    now = time.time()
    meta_set("bets_top_cleared_at", now)
    deleted = 0
    if hard == 1:
        try:
            cnt_before = (query_one("SELECT COUNT(*) AS c FROM bets") or {}).get("c", 0)
            execute("DELETE FROM bets")
            deleted = cnt_before
        except Exception as e:
            raise HTTPException(500, f"delete failed: {e}")
    return {"ok": True, "cleared_at": now, "hard": bool(hard), "deleted": deleted}


@app.post("/api/bets_to_rains/clear")
async def clear_bets_to_rains():
    meta_set("bets_to_rains_cleared_at", time.time())
    return {"ok": True, "cleared_at": time.time()}


@app.post("/api/intent")
async def add_intent(body: IntentIn, x_api_key: str = Header(default="")):
    if x_api_key != API_KEY:
        raise HTTPException(401, "Invalid API key")
    ts_val = body.ts if body.ts else time.time()
    dup = query_one("SELECT id FROM intents WHERE chat=? AND sender=? AND text=? AND ABS(ts-?)<60 LIMIT 1",
                    (body.chat, body.sender, body.text, ts_val))
    if dup:
        return {"ok": True, "dup": True}
    p = ph()
    execute(f"""INSERT INTO intents (ts,chat,sender,text,percent,matched_rain_id,matched_delta_sec)
                VALUES ({p},{p},{p},{p},{p},NULL,NULL)""",
            (ts_val, body.chat, body.sender, body.text, body.percent or 100))
    rain = query_one("""SELECT id, ts FROM events WHERE type='rain' AND chat=? AND sender LIKE ?
                        AND ts IS NOT NULL AND ts>=? AND ts<=? ORDER BY ts ASC LIMIT 1""",
                     (body.chat, f"%{body.sender}%", ts_val, ts_val + 300))
    if rain:
        delta = int(rain["ts"] - ts_val)
        execute("""UPDATE intents SET matched_rain_id=?, matched_delta_sec=?
                   WHERE chat=? AND sender=? AND text=? AND ABS(ts-?)<60""",
                (rain["id"], delta, body.chat, body.sender, body.text, ts_val))
    return {"ok": True}


@app.post("/api/players_batch")
async def add_players_batch(body: PlayersBatchIn, x_api_key: str = Header(default="")):
    if x_api_key != API_KEY:
        raise HTTPException(401, "Invalid API key")
    now = time.time()
    p = ph()
    for pl in body.players:
        name = (pl.name or "").strip()
        if not name or name == "?":
            continue
        msgs = max(0, int(pl.msgs))
        last_seen = float(pl.last_seen) or now
        if USE_POSTGRES:
            execute(f"""INSERT INTO chat_players (chat,username,msg_count,first_seen,last_seen)
                        VALUES ({p},{p},{p},{p},{p})
                        ON CONFLICT (chat,username) DO UPDATE
                        SET msg_count=chat_players.msg_count+EXCLUDED.msg_count,
                            last_seen=GREATEST(chat_players.last_seen,EXCLUDED.last_seen)""",
                    (body.chat, name, msgs, last_seen, last_seen))
        else:
            execute(f"""INSERT INTO chat_players (chat,username,msg_count,first_seen,last_seen)
                        VALUES ({p},{p},{p},{p},{p})
                        ON CONFLICT (chat,username) DO UPDATE
                        SET msg_count=msg_count+excluded.msg_count,
                            last_seen=MAX(last_seen,excluded.last_seen)""",
                    (body.chat, name, msgs, last_seen, last_seen))
    if body.ticks:
        for t in body.ticks:
            un = (t.username or "").strip()
            if not un or un == "?":
                continue
            try:
                ts_val = float(t.ts or 0)
            except Exception:
                continue
            if ts_val <= 0:
                continue
            if USE_POSTGRES:
                execute(f"""INSERT INTO player_message_ticks (username,chat,ts)
                            VALUES ({p},{p},{p}) ON CONFLICT DO NOTHING""",
                        (un, body.chat, ts_val))
            else:
                execute(f"""INSERT OR IGNORE INTO player_message_ticks (username,chat,ts)
                            VALUES ({p},{p},{p})""",
                        (un, body.chat, ts_val))
    return {"ok": True}


@app.get("/api/player/activity")
async def api_player_activity(username: str, windows: str = "15,25,60,1440"):
    un = (username or "").strip()
    if not un:
        raise HTTPException(400, "username required")
    try:
        wins = sorted({int(w) for w in windows.split(",") if w.strip()})
    except Exception:
        wins = [15, 25, 60, 1440]
    max_win = max(wins) if wins else 1440
    since = time.time() - max_win * 60
    all_chat_rows = query_all("SELECT chat, msg_count, last_seen FROM chat_players WHERE username=?",
                              (un,))
    all_chats = {
        r["chat"]: {"msg_count": r.get("msg_count", 0) or 0,
                    "last_seen_total": r.get("last_seen") or 0}
        for r in all_chat_rows
    }
    tick_rows = query_all("""SELECT chat, ts FROM player_message_ticks
                             WHERE username=? AND ts>=? ORDER BY ts DESC""",
                          (un, since))
    now = time.time()
    per_chat = {
        c: {"counts": {w: 0 for w in wins}, "last_ts": 0.0,
            "total_msgs": m["msg_count"] or 0, "last_seen_total": m["last_seen_total"] or 0}
        for c, m in all_chats.items()
    }
    for r in tick_rows:
        c = r["chat"]
        ts = float(r["ts"] or 0)
        if c not in per_chat:
            per_chat[c] = {"counts": {w: 0 for w in wins}, "last_ts": 0.0,
                           "total_msgs": 0, "last_seen_total": 0}
        for w in wins:
            if now - ts <= w * 60:
                per_chat[c]["counts"][w] += 1
        if ts > per_chat[c]["last_ts"]:
            per_chat[c]["last_ts"] = ts
    total = {w: 0 for w in wins}
    last_ts = 0.0
    for c, v in per_chat.items():
        for w in wins:
            total[w] += v["counts"][w]
        if v["last_ts"] > last_ts:
            last_ts = v["last_ts"]
    chats = []
    for c, v in per_chat.items():
        eff = max(v["last_ts"] or 0, v["last_seen_total"] or 0)
        chats.append({"chat": c, "counts": v["counts"], "last_ts": eff,
                      "minutes_ago": int((now - eff) / 60) if eff else None,
                      "total_msgs": v["total_msgs"] or 0})
    chats.sort(key=lambda x: (-sum(x["counts"].values()), -x["last_ts"]))
    gl = max([c["last_ts"] for c in chats], default=0)
    return {"username": un, "windows": wins, "total": total, "last_ts": gl,
            "minutes_ago_total": int((now - gl) / 60) if gl else None,
            "chats": chats, "now": now, "known_chats": len(chats)}


@app.get("/api/user/{nickname}/rains_received")
async def user_rains_received(nickname: str, days: int = 30, limit: int = 500):
    since = time.time() - days * 86400
    rows = query_all("""SELECT id, ts, chat, sender, receivers, amount_text, amount_rub,
                        amount_usd, crypto_symbol FROM events
                        WHERE type='rain' AND ts >= ? ORDER BY ts DESC LIMIT 5000""",
                     (since,))
    out = []
    total_rub = 0.0
    total_usd = 0.0
    for r in rows:
        try:
            recs = json.loads(r.get("receivers") or "[]")
        except Exception:
            recs = []
        if nickname not in recs:
            continue
        n = len(recs) or 1
        per_rub = (r.get("amount_rub") or 0) / n
        per_usd = (r.get("amount_usd") or 0) / n
        total_rub += per_rub
        total_usd += per_usd
        out.append({"id": r["id"], "ts": r["ts"], "chat": r["chat"], "sender": r["sender"],
                    "amount_text": r["amount_text"],
                    "crypto_symbol": r.get("crypto_symbol", ""),
                    "share_rub": round(per_rub, 2), "share_usd": round(per_usd, 2),
                    "receivers_count": n})
        if len(out) >= limit:
            break
    return {"nickname": nickname, "wins": len(out),
            "total_rub": round(total_rub, 2), "total_usd": round(total_usd, 2),
            "items": out, "days": days}


@app.post("/api/status")
async def post_status(payload: StatusPayload, x_api_key: str = Header(default="")):
    if x_api_key != API_KEY:
        raise HTTPException(401, "Invalid API key")
    now = time.time()
    with _state_lock:
        for inst in payload.instances:
            INSTANCE_STATES[inst.iid] = {**inst.dict(), "server_received_at": now}
        for iid in payload.screenshot_requests:
            if iid not in SCREENSHOT_REQUESTS:
                SCREENSHOT_REQUESTS[iid] = {"requested_at": now, "done": False}
    return {"ok": True}


@app.get("/api/screenshot_req")
async def get_screenshot_req(x_api_key: str = Header(default="")):
    if x_api_key != API_KEY:
        raise HTTPException(401, "Invalid API key")
    with _state_lock:
        pending = [{"iid": iid, "requested_at": r.get("requested_at", 0)}
                   for iid, r in SCREENSHOT_REQUESTS.items() if not r.get("done")]
    return {"requests": pending}


@app.post("/api/screenshot_upload")
async def upload_screenshot(iid: int = Form(...), file: UploadFile = File(...),
                            x_api_key: str = Header(default="")):
    if x_api_key != API_KEY:
        raise HTTPException(401, "Invalid API key")
    path = os.path.join(SCREENSHOT_DIR, f"inst{iid}_latest.png")
    content = await file.read()
    with open(path, "wb") as f:
        f.write(content)
    with _state_lock:
        prev = SCREENSHOT_REQUESTS.get(iid, {})
        SCREENSHOT_REQUESTS[iid] = {"requested_at": prev.get("requested_at", time.time()),
                                     "done": True, "path": path,
                                     "uploaded_at": time.time()}
    return {"ok": True}


# ==================== CONTROL / NGROK / MODERATORS / LOG FILTER ====================
@app.get("/api/control/get")
async def control_get():
    v = meta_get("control_action") or ""
    ts = float(meta_get("control_ts") or 0)
    return {"action": v, "ts": ts}


@app.post("/api/control/set")
async def control_set(body: ControlSet):
    if body.action not in ("start", "stop", "restart"):
        raise HTTPException(400, "invalid action")
    meta_set("control_action", body.action)
    meta_set("control_ts", str(time.time()))
    return {"ok": True, "action": body.action}


@app.post("/api/ngrok_url")
async def post_ngrok_url(body: NgrokUrlIn, x_api_key: str = Header(default="")):
    if x_api_key != API_KEY:
        raise HTTPException(401, "Invalid API key")
    meta_set("ngrok_url", body.url)
    meta_set("ngrok_url_ts", str(body.ts or time.time()))
    return {"ok": True}


@app.get("/api/ngrok_url")
async def get_ngrok_url():
    return {"url": meta_get("ngrok_url") or "",
            "ts": float(meta_get("ngrok_url_ts") or 0)}


@app.get("/api/log_filter")
async def get_log_filter():
    v = meta_get("log_filter")
    if not v:
        return {"categories": DEFAULT_LOG_FILTER}
    try:
        return {"categories": json.loads(v)}
    except Exception:
        return {"categories": DEFAULT_LOG_FILTER}


@app.post("/api/log_filter")
async def set_log_filter(body: LogFilterIn):
    meta_set("log_filter", json.dumps(body.categories, ensure_ascii=False))
    return {"ok": True, "categories": body.categories}


@app.post("/api/moderator_message")
async def post_moderator_message(body: ModeratorMessageIn, x_api_key: str = Header(default="")):
    if x_api_key != API_KEY:
        raise HTTPException(401, "Invalid API key")
    ts_val = body.ts or time.time()
    dup = query_one("""SELECT id FROM moderator_messages
                       WHERE username=? AND chat=? AND text=?
                       AND ABS(ts-?) < 300 LIMIT 1""",
                    (body.username, body.chat, body.text, ts_val))
    if dup:
        return {"ok": True, "dup": True}
    p = ph()
    execute(f"""INSERT INTO moderator_messages (ts,username,chat,text)
                VALUES ({p},{p},{p},{p})""",
            (ts_val, body.username, body.chat, body.text))
    row = query_one("SELECT * FROM moderators WHERE username=?", (body.username,))
    if row:
        try:
            chats = json.loads(row.get("chats") or "{}")
        except Exception:
            chats = {}
        chats[body.chat] = chats.get(body.chat, 0) + 1
        execute(f"UPDATE moderators SET last_seen={p}, chats={p} WHERE username={p}",
                (ts_val, json.dumps(chats), body.username))
    else:
        chats = {body.chat: 1}
        execute(f"""INSERT INTO moderators (username,first_seen,last_seen,chats)
                    VALUES ({p},{p},{p},{p})""",
                (body.username, ts_val, ts_val, json.dumps(chats)))
    return {"ok": True}


@app.get("/api/moderators")
async def get_moderators():
    rows = query_all("SELECT * FROM moderators ORDER BY last_seen DESC LIMIT 500")
    for r in rows:
        try:
            r["chats"] = json.loads(r.get("chats") or "{}")
        except Exception:
            r["chats"] = {}
    return {"moderators": rows}


@app.get("/api/moderator_messages")
async def get_moderator_messages(username: Optional[str] = None,
                                  chat: Optional[str] = None,
                                  limit: int = 200):
    q = "SELECT * FROM moderator_messages WHERE 1=1"
    p = []
    if username:
        q += " AND username=?"
        p.append(username)
    if chat:
        q += " AND chat=?"
        p.append(chat)
    q += " ORDER BY ts DESC LIMIT ?"
    p.append(limit)
    return {"messages": query_all(q, tuple(p))}


# ==================== WATCHLIST / PROFILES ====================
DEFAULT_WATCHLIST = ["Fab", "Fab434", "fab434", MY_NICK]


@app.get("/api/watchlist")
async def get_watchlist():
    v = meta_get("watchlist")
    if not v:
        return {"players": DEFAULT_WATCHLIST}
    try:
        return {"players": json.loads(v)}
    except Exception:
        return {"players": DEFAULT_WATCHLIST}


@app.post("/api/watchlist/add")
async def watchlist_add(body: dict):
    name = (body.get("name") or "").strip()
    if not name:
        raise HTTPException(400, "name required")
    v = meta_get("watchlist")
    try:
        players = json.loads(v) if v else list(DEFAULT_WATCHLIST)
    except Exception:
        players = list(DEFAULT_WATCHLIST)
    if name not in players:
        players.append(name)
    meta_set("watchlist", json.dumps(players, ensure_ascii=False))
    return {"ok": True, "players": players}


@app.post("/api/watchlist/remove")
async def watchlist_remove(body: dict):
    name = (body.get("name") or "").strip()
    v = meta_get("watchlist")
    try:
        players = json.loads(v) if v else list(DEFAULT_WATCHLIST)
    except Exception:
        players = list(DEFAULT_WATCHLIST)
    players = [p for p in players if p != name]
    meta_set("watchlist", json.dumps(players, ensure_ascii=False))
    return {"ok": True, "players": players}


@app.get("/api/profile/{username}")
async def get_profile(username: str):
    p = query_one("SELECT * FROM player_profiles WHERE username=?", (username,))
    if not p:
        return {"profile": {"username": username, "lang": "ENGLISH",
                            "description": "", "is_top": False}}
    return {"profile": p}


@app.post("/api/profile/{username}")
async def upsert_profile(username: str, body: dict):
    lang = (body.get("lang") or "ENGLISH").strip()
    desc = (body.get("description") or "").strip()
    is_top = 1 if body.get("is_top") else 0
    p = ph()
    if USE_POSTGRES:
        execute(f"""INSERT INTO player_profiles (username,lang,description,is_top,updated_at)
                    VALUES ({p},{p},{p},{p},{p}) ON CONFLICT (username) DO UPDATE
                    SET lang=EXCLUDED.lang,description=EXCLUDED.description,
                        is_top=EXCLUDED.is_top,updated_at=EXCLUDED.updated_at""",
                (username, lang, desc, is_top, time.time()))
    else:
        execute(f"""INSERT INTO player_profiles (username,lang,description,is_top,updated_at)
                    VALUES ({p},{p},{p},{p},{p}) ON CONFLICT (username) DO UPDATE
                    SET lang=excluded.lang,description=excluded.description,
                        is_top=excluded.is_top,updated_at=excluded.updated_at""",
                (username, lang, desc, is_top, time.time()))
    return {"ok": True}


# ==================== ПРОМОКОДЫ ====================
def fetch_promo_channel():
    try:
        r = requests.get(f"https://t.me/s/{PROMO_CHANNEL}",
                         headers={"User-Agent": "Mozilla/5.0"}, timeout=15)
        if r.status_code != 200:
            return []
        soup = BeautifulSoup(r.text, "html.parser")
        messages = []
        for wrap in soup.select(".tgme_widget_message_wrap"):
            msg_el = wrap.select_one(".tgme_widget_message")
            if not msg_el:
                continue
            data_post = msg_el.get("data-post", "")
            if not data_post:
                continue
            time_el = wrap.select_one("time")
            msg_date = time_el.get("datetime", "") if time_el else ""
            text_el = wrap.select_one(".tgme_widget_message_text")
            text = text_el.get_text("\n", strip=True) if text_el else ""
            if not text:
                continue
            messages.append({"tg_msg_id": data_post, "text": text,
                             "link": f"https://t.me/{data_post}",
                             "message_date": msg_date})
        return messages
    except Exception:
        return []


def promo_parser_thread():
    time.sleep(10)
    while True:
        try:
            messages = fetch_promo_channel()
            for m in messages:
                if query_one("SELECT id FROM promocodes WHERE tg_msg_id=?", (m["tg_msg_id"],)):
                    continue
                execute(f"""INSERT INTO promocodes
                            (tg_msg_id,text,link,message_date,discovered_at,used)
                            VALUES ({ph()},{ph()},{ph()},{ph()},{ph()},0)""",
                        (m["tg_msg_id"], m["text"], m["link"], m["message_date"],
                         time.time()))
        except Exception:
            pass
        time.sleep(PROMO_PARSE_INTERVAL)


@app.get("/api/promocodes")
async def api_promocodes(only_unused: Optional[int] = None,
                         only_used: Optional[int] = None,
                         limit: int = 200):
    q = "SELECT * FROM promocodes WHERE 1=1"
    p = []
    if only_unused == 1:
        q += " AND used = 0"
    if only_used == 1:
        q += " AND used = 1"
    q += " ORDER BY discovered_at DESC LIMIT ?"
    p.append(limit)
    return {"promos": query_all(q, tuple(p))}


@app.post("/api/promo/toggle/{promo_id}")
async def api_promo_toggle(promo_id: int):
    row = query_one("SELECT id, used FROM promocodes WHERE id=?", (promo_id,))
    if not row:
        raise HTTPException(404, "not found")
    nv = 0 if row["used"] else 1
    execute("UPDATE promocodes SET used=?, used_at=? WHERE id=?",
            (nv, time.time() if nv else None, promo_id))
    return {"ok": True, "used": nv}


@app.post("/api/promo/to_vault/{promo_id}")
async def promo_to_vault(promo_id: int, body: dict = None):
    row = query_one("SELECT * FROM promocodes WHERE id=?", (promo_id,))
    if not row:
        raise HTTPException(404, "not found")
    body = body or {}
    amount_usd = float(body.get("amount_usd") or 0)
    currency = (body.get("currency") or "USD").upper()
    execute("UPDATE promocodes SET used_for_vault=1, vault_amount_usd=? WHERE id=?",
            (amount_usd, promo_id))
    execute(f"""INSERT INTO vault_transactions (ts,type,currency,amount,amount_usd,description)
                VALUES ({ph()},{ph()},{ph()},{ph()},{ph()},{ph()})""",
            (time.time(), "promo_code", currency, amount_usd, amount_usd,
             f"Промокод: {row['text'][:100]}"))
    return {"ok": True}


@app.get("/api/promocodes/stats")
async def api_promocodes_stats():
    total = (query_one("SELECT COUNT(*) AS c FROM promocodes") or {}).get("c", 0)
    used = (query_one("SELECT COUNT(*) AS c FROM promocodes WHERE used=1") or {}).get("c", 0)
    return {"total": total, "used": used, "unused": total - used}


@app.post("/api/promo/test")
async def api_promo_test():
    messages = fetch_promo_channel()
    new_count = 0
    for m in messages:
        if query_one("SELECT id FROM promocodes WHERE tg_msg_id=?", (m["tg_msg_id"],)):
            continue
        execute(f"""INSERT INTO promocodes
                    (tg_msg_id,text,link,message_date,discovered_at,used)
                    VALUES ({ph()},{ph()},{ph()},{ph()},{ph()},0)""",
                (m["tg_msg_id"], m["text"], m["link"], m["message_date"], time.time()))
        new_count += 1
    return {"ok": True, "found": len(messages), "new": new_count}


# ==================== VAULT ====================
def _balance_upsert(table, cur, amt, amt_usd):
    p = ph()
    if USE_POSTGRES:
        execute(f"""INSERT INTO {table} (currency,amount,amount_usd,updated_at)
                    VALUES ({p},{p},{p},{p})
                    ON CONFLICT (currency) DO UPDATE
                    SET amount=EXCLUDED.amount,amount_usd=EXCLUDED.amount_usd,
                        updated_at=EXCLUDED.updated_at""",
                (cur, amt, amt_usd, time.time()))
    else:
        execute(f"""INSERT INTO {table} (currency,amount,amount_usd,updated_at)
                    VALUES ({p},{p},{p},{p})
                    ON CONFLICT (currency) DO UPDATE
                    SET amount=excluded.amount,amount_usd=excluded.amount_usd,
                        updated_at=excluded.updated_at""",
                (cur, amt, amt_usd, time.time()))


def _balance_get(table, cur):
    return query_one(f"SELECT amount, amount_usd FROM {table} WHERE currency=?", (cur,))


def _balance_add(table, cur, delta_amt, delta_usd):
    row = _balance_get(table, cur)
    new_amt = max(0.0, (float(row["amount"]) if row else 0.0) + delta_amt)
    new_usd = max(0.0, (float(row["amount_usd"]) if row else 0.0) + delta_usd)
    _balance_upsert(table, cur, new_amt, new_usd)


def _calc_usd(cur, amount):
    cur = (cur or "").upper().strip()
    if cur in ("USD", "USDT", "USDC", "DAI", "CASH"):
        return amount
    rate = server_rate_usd(cur)
    return amount * rate if rate else 0.0


@app.get("/api/vault")
async def get_vault():
    account = query_all("SELECT * FROM account_balances ORDER BY currency")
    vault = query_all("SELECT * FROM vault_balances ORDER BY currency")
    wager = query_one("SELECT * FROM wager_state WHERE id=1") or {"amount_required": 0, "deadline_ts": 0}
    goal_usd = float(vault_meta_get("goal_usd") or 0)
    goal_text = vault_meta_get("goal_text") or ""
    now = time.time()
    week_ago = now - 7 * 86400
    income_week = (query_one("""SELECT COALESCE(SUM(amount_usd),0) AS s FROM vault_transactions
                                WHERE ts>=? AND type='income'""", (week_ago,)) or {"s": 0})["s"] or 0
    wager_week = (query_one("""SELECT COALESCE(SUM(amount_usd),0) AS s FROM vault_transactions
                               WHERE ts>=? AND type='wager_spend'""", (week_ago,)) or {"s": 0})["s"] or 0
    vault_add_week = (query_one("""SELECT COALESCE(SUM(amount_usd),0) AS s FROM vault_transactions
                                   WHERE ts>=? AND type='vault_add'""", (week_ago,)) or {"s": 0})["s"] or 0
    total_account = sum(float(b.get("amount_usd") or 0) for b in account)
    total_vault = sum(float(b.get("amount_usd") or 0) for b in vault)
    days_left = max(0, int((wager.get("deadline_ts", 0) - now) / 86400)) if wager.get("deadline_ts") else None
    hours_left = max(0, int((wager.get("deadline_ts", 0) - now) / 3600)) if wager.get("deadline_ts") else None
    return {
        "account": account, "vault": vault,
        "total_account_usd": round(total_account, 2),
        "total_vault_usd": round(total_vault, 2),
        "total_usd": round(total_account + total_vault, 2),
        "wager": {"amount_required": wager.get("amount_required", 0),
                  "deadline_ts": wager.get("deadline_ts", 0),
                  "days_left": days_left, "hours_left": hours_left,
                  "is_expired": (wager.get("deadline_ts", 0) > 0 and wager.get("deadline_ts", 0) < now)},
        "goal_usd": goal_usd, "goal_text": goal_text,
        "week": {"income": round(income_week, 2), "wager_spend": round(wager_week, 2),
                 "vault_add": round(vault_add_week, 2)},
        "currencies": CURRENCIES, "now": now,
    }


@app.post("/api/account/balance")
async def set_account_balance(body: BalanceIn):
    cur = body.currency.upper().strip()
    if not cur:
        raise HTTPException(400, "currency required")
    amt = float(body.amount or 0)
    amt_usd = body.amount_usd if body.amount_usd is not None else _calc_usd(cur, amt)
    _balance_upsert("account_balances", cur, amt, amt_usd)
    return {"ok": True}


@app.post("/api/account/balance/delete")
async def delete_account_balance(body: BalanceDeleteIn):
    cur = (body.currency or "").upper().strip()
    if not cur:
        raise HTTPException(400, "currency required")
    execute("DELETE FROM account_balances WHERE currency=?", (cur,))
    return {"ok": True}


@app.post("/api/vault/balance")
async def set_vault_balance(body: BalanceIn):
    cur = body.currency.upper().strip()
    if not cur:
        raise HTTPException(400, "currency required")
    amt = float(body.amount or 0)
    amt_usd = body.amount_usd if body.amount_usd is not None else _calc_usd(cur, amt)
    _balance_upsert("vault_balances", cur, amt, amt_usd)
    return {"ok": True}


@app.post("/api/vault/balance/delete")
async def delete_vault_balance(body: BalanceDeleteIn):
    cur = (body.currency or "").upper().strip()
    if not cur:
        raise HTTPException(400, "currency required")
    execute("DELETE FROM vault_balances WHERE currency=?", (cur,))
    return {"ok": True}


@app.post("/api/transfer")
async def transfer(body: TransferIn):
    direction = body.direction.strip()
    cur = body.currency.upper().strip()
    amt = float(body.amount or 0)
    if direction not in ("to_vault", "to_account"):
        raise HTTPException(400, "direction must be to_vault or to_account")
    if amt <= 0:
        raise HTTPException(400, "amount must be > 0")
    if not cur:
        raise HTTPException(400, "currency required")
    usd = _calc_usd(cur, amt)
    if direction == "to_vault":
        row = _balance_get("account_balances", cur)
        if not row or float(row["amount"]) < amt:
            raise HTTPException(400, f"Недостаточно {cur} на аккаунте")
        _balance_add("account_balances", cur, -amt, -usd)
        _balance_add("vault_balances", cur, amt, usd)
        execute(f"""INSERT INTO vault_transactions (ts,type,currency,amount,amount_usd,description)
                    VALUES ({ph()},'vault_add',{ph()},{ph()},{ph()},{ph()})""",
                (time.time(), cur, amt, usd, f"Перевод в хранилище {amt} {cur}"))
    else:
        row = _balance_get("vault_balances", cur)
        if not row or float(row["amount"]) < amt:
            raise HTTPException(400, f"Недостаточно {cur} в хранилище")
        _balance_add("vault_balances", cur, -amt, -usd)
        _balance_add("account_balances", cur, amt, usd)
        execute(f"""INSERT INTO vault_transactions (ts,type,currency,amount,amount_usd,description)
                    VALUES ({ph()},'vault_take',{ph()},{ph()},{ph()},{ph()})""",
                (time.time(), cur, amt, usd, f"Перевод из хранилища {amt} {cur}"))
    return {"ok": True}


@app.post("/api/vault/transaction")
async def add_vault_transaction(body: VaultTransactionIn):
    typ = body.type.strip()
    if typ not in ("income", "wager_spend", "vault_add", "vault_take", "promo_code"):
        raise HTTPException(400, "invalid type")
    ts_val = body.ts if body.ts else time.time()
    cur = (body.currency or "USDT").upper().strip()
    amount = float(body.amount or 0)
    amount_usd = float(body.amount_usd or 0)
    if amount_usd == 0 and amount != 0:
        amount_usd = _calc_usd(cur, amount)
    execute(f"""INSERT INTO vault_transactions (ts,type,currency,amount,amount_usd,description)
                VALUES ({ph()},{ph()},{ph()},{ph()},{ph()},{ph()})""",
            (ts_val, typ, cur, amount, amount_usd, body.description or ""))
    if typ == "vault_add" and amount:
        _balance_add("vault_balances", cur, amount, amount_usd)
        _balance_add("account_balances", cur, -amount, -amount_usd)
    if typ == "vault_take" and amount:
        _balance_add("vault_balances", cur, -amount, -amount_usd)
        _balance_add("account_balances", cur, amount, amount_usd)
    if typ == "income" and amount:
        _balance_add("account_balances", cur, amount, amount_usd)
    return {"ok": True}


@app.get("/api/vault/transactions")
async def list_vault_transactions(days: int = 30, type: Optional[str] = None, limit: int = 500):
    since = time.time() - days * 86400
    q = "SELECT * FROM vault_transactions WHERE ts>=?"
    p = [since]
    if type:
        q += " AND type=?"
        p.append(type)
    q += " ORDER BY ts DESC LIMIT ?"
    p.append(limit)
    return {"transactions": query_all(q, tuple(p))}


@app.post("/api/vault/wager")
async def set_wager(body: VaultWagerIn):
    p = ph()
    if USE_POSTGRES:
        execute(f"""INSERT INTO wager_state (id,amount_required,deadline_ts,updated_at)
                    VALUES (1,{p},{p},{p}) ON CONFLICT (id) DO UPDATE
                    SET amount_required=EXCLUDED.amount_required,
                        deadline_ts=EXCLUDED.deadline_ts,
                        updated_at=EXCLUDED.updated_at""",
                (body.amount_required, body.deadline_ts, time.time()))
    else:
        execute(f"""INSERT INTO wager_state (id,amount_required,deadline_ts,updated_at)
                    VALUES (1,{p},{p},{p}) ON CONFLICT (id) DO UPDATE
                    SET amount_required=excluded.amount_required,
                        deadline_ts=excluded.deadline_ts,
                        updated_at=excluded.updated_at""",
                (body.amount_required, body.deadline_ts, time.time()))
    return {"ok": True}


@app.post("/api/vault/goal")
async def set_goal(body: VaultGoalIn):
    vault_meta_set("goal_usd", body.goal_usd)
    vault_meta_set("goal_text", body.goal_text)
    return {"ok": True}


# ==================== AI ====================
@app.get("/api/chat/pending")
async def chat_pending(x_api_key: str = Header(default="")):
    if x_api_key != API_KEY:
        raise HTTPException(401, "Invalid API key")
    task = query_one("SELECT id, chat FROM chat_tasks WHERE status='pending' ORDER BY created_at ASC LIMIT 1")
    if not task:
        return {"task": None}
    execute("UPDATE chat_tasks SET status='reading' WHERE id=?", (task["id"],))
    return {"task": task}


@app.post("/api/chat/complete")
async def chat_complete(body: ChatTaskCompleteIn, x_api_key: str = Header(default="")):
    if x_api_key != API_KEY:
        raise HTTPException(401, "Invalid API key")
    task = query_one("SELECT * FROM chat_tasks WHERE id=?", (body.task_id,))
    if not task:
        raise HTTPException(404, "task not found")
    if body.error:
        execute("UPDATE chat_tasks SET status='failed', error=?, completed_at=? WHERE id=?",
                (body.error, time.time(), body.task_id))
        return {"ok": True, "status": "failed"}
    execute("UPDATE chat_tasks SET messages=? WHERE id=?",
            (json.dumps(body.messages, ensure_ascii=False), body.task_id))
    threading.Thread(target=_process_chat_task_groq,
                     args=(body.task_id, task["chat"], body.messages), daemon=True).start()
    return {"ok": True, "status": "processing"}


def _process_chat_task_groq(task_id, chat_name, messages):
    if not GROQ_API_KEY or not GROQ_API_KEY.startswith("gsk_"):
        execute("UPDATE chat_tasks SET status='failed', error='GROQ_API_KEY не задан', completed_at=? WHERE id=?",
                (time.time(), task_id))
        return
    lang = LANG_MAP.get(chat_name, chat_name.lower())
    lines = [f"{m.get('user', '?')}: {m.get('text', '')}"
             for m in (messages or [])[-50:] if m.get("text")]
    context = "\n".join(lines)
    sp = (f"Ты носитель {lang} и активный участник игрового чата казино. "
          f"Дан контекст из последних 50 сообщений. Предложи 5 коротких живых сообщений на {lang} (до 100 символов) "
          f"и 5 нейтральных фраз. Ответь СТРОГО JSON: "
          f'{{"topic": "...", "messages": ["..."], "phrases": ["..."]}}')
    models = [GROQ_MODEL, "openai/gpt-oss-20b", "openai/gpt-oss-120b"]
    content = None
    last_error = ""
    for m in models:
        try:
            r = requests.post("https://api.groq.com/openai/v1/chat/completions",
                              headers={"Authorization": f"Bearer {GROQ_API_KEY}",
                                       "Content-Type": "application/json"},
                              json={"model": m,
                                    "messages": [{"role": "system", "content": sp},
                                                 {"role": "user", "content": context}],
                                    "max_tokens": 800, "temperature": 0.8},
                              timeout=30)
            if r.status_code == 200:
                content = r.json().get("choices", [{}])[0].get("message", {}).get("content", "").strip()
                break
            elif r.status_code == 404:
                last_error = f"{m} not found"
                continue
            else:
                last_error = f"HTTP {r.status_code}"
                break
        except Exception as e:
            last_error = f"{type(e).__name__}"
    if not content:
        execute("UPDATE chat_tasks SET status='failed', error=?, completed_at=? WHERE id=?",
                (last_error, time.time(), task_id))
        return
    mm = re.search(r'\{[\s\S]*\}', content)
    if not mm:
        execute("UPDATE chat_tasks SET status='failed', error='не-JSON', completed_at=? WHERE id=?",
                (time.time(), task_id))
        return
    try:
        data = json.loads(mm.group(0))
    except Exception as e:
        execute("UPDATE chat_tasks SET status='failed', error=?, completed_at=? WHERE id=?",
                (f"JSON: {e}", time.time(), task_id))
        return
    result = {"topic": data.get("topic", ""),
              "messages": data.get("messages", [])[:5],
              "phrases": data.get("phrases", [])[:5]}
    execute("UPDATE chat_tasks SET status='completed', result=?, completed_at=? WHERE id=?",
            (json.dumps(result, ensure_ascii=False), time.time(), task_id))


@app.post("/api/chat/request")
async def chat_request(body: ChatTaskIn):
    if not body.chat:
        raise HTTPException(400, "chat required")
    execute(f"INSERT INTO chat_tasks (chat,status,created_at) VALUES ({ph()},'pending',{ph()})",
            (body.chat, time.time()))
    task = query_one("SELECT id FROM chat_tasks WHERE chat=? AND status='pending' ORDER BY id DESC LIMIT 1",
                     (body.chat,))
    return {"ok": True, "task_id": task["id"] if task else None}


@app.get("/api/chat/task/{task_id}")
async def chat_task_get(task_id: int):
    task = query_one("SELECT * FROM chat_tasks WHERE id=?", (task_id,))
    if not task:
        raise HTTPException(404, "not found")
    result = None
    if task.get("result"):
        try:
            result = json.loads(task["result"])
        except Exception:
            pass
    return {"id": task["id"], "chat": task["chat"], "status": task["status"],
            "result": result, "error": task.get("error"),
            "created_at": task.get("created_at"),
            "completed_at": task.get("completed_at")}


@app.get("/api/chat/latest")
async def chat_latest(chat: Optional[str] = None):
    q = "SELECT * FROM chat_tasks WHERE status='completed'"
    p = []
    if chat:
        q += " AND chat=?"
        p.append(chat)
    q += " ORDER BY completed_at DESC LIMIT 1"
    task = query_one(q, tuple(p))
    if not task:
        return {"task": None}
    result = None
    if task.get("result"):
        try:
            result = json.loads(task["result"])
        except Exception:
            pass
    return {"task": {"id": task["id"], "chat": task["chat"],
                     "result": result, "completed_at": task["completed_at"]}}


# ==================== ACTIVITY / RAIN / DOCX ====================
@app.get("/api/activity/dashboard")
async def api_activity_dashboard(minutes: int = 60):
    now = time.time()
    wl_raw = meta_get("watchlist")
    try:
        watchlist = json.loads(wl_raw) if wl_raw else DEFAULT_WATCHLIST
    except Exception:
        watchlist = DEFAULT_WATCHLIST
    if MY_NICK not in watchlist:
        watchlist = list(watchlist) + [MY_NICK]

    rows = []
    if watchlist:
        phs = ",".join(["?"] * len(watchlist))
        rows = query_all(f"SELECT chat, username, msg_count, last_seen FROM chat_players WHERE username IN ({phs})",
                         tuple(watchlist))

    players = {name: {"total_msgs": 0, "last_seen": 0.0, "chats": {}} for name in watchlist}
    for r in rows:
        un = r.get("username")
        if un not in players:
            continue
        chat = r.get("chat")
        last_seen = float(r.get("last_seen") or 0)
        cnt = int(r.get("msg_count") or 0)
        minutes_ago = int((now - last_seen) / 60) if last_seen > 0 else None
        players[un]["chats"][chat] = {
            "last_seen": last_seen,
            "minutes_ago": minutes_ago,
            "active_60m": (last_seen > 0 and (now - last_seen) <= 3600),
            "active_N": (last_seen > 0 and (now - last_seen) <= minutes * 60),
            "msg_count": cnt,
        }
        players[un]["total_msgs"] += cnt
        if last_seen > players[un]["last_seen"]:
            players[un]["last_seen"] = last_seen
    return {"watchlist": watchlist, "players": players, "now": now, "window_min": minutes}


@app.get("/api/rain_activity")
async def api_rain_activity(hours: int = 24, limit: int = 100, chat: Optional[str] = None):
    since = time.time() - hours * 3600
    q = """SELECT id, ts, chat, sender, receivers, receivers_info,
                  amount_text, amount_rub, amount_usd, crypto_symbol
           FROM events WHERE type='rain' AND ts >= ?"""
    p = [since]
    if chat:
        q += " AND chat=?"
        p.append(chat)
    q += " ORDER BY ts DESC LIMIT ?"
    p.append(limit)
    rows = query_all(q, tuple(p))

    out = []
    for r in rows:
        try:
            recs = json.loads(r.get("receivers") or "[]")
        except Exception:
            recs = []
        try:
            rinfo = json.loads(r.get("receivers_info") or "[]")
        except Exception:
            rinfo = []
        rinfo_sorted = sorted(rinfo, key=lambda x: (-(x.get("total_msgs") or 0)))
        top20 = rinfo_sorted[:20]
        won = MY_NICK in recs

        my_rec = next((x for x in rinfo if x.get("name") == MY_NICK), None)
        my_minutes = my_rec.get("minutes_ago") if my_rec else None
        my_was_active = bool(my_rec and my_rec.get("wrote_in_25min"))

        # Фолбэк — ищем последнее сообщение Merami ДО времени дождя
        if my_minutes is None:
            rain_ts = float(r.get("ts") or 0)
            if rain_ts > 0:
                try:
                    tick = query_one(
                        """SELECT MAX(ts) AS last_before FROM player_message_ticks
                           WHERE username=? AND chat=? AND ts <= ?""",
                        (MY_NICK, r["chat"], rain_ts)
                    )
                    if tick and tick.get("last_before"):
                        tick_ts = float(tick["last_before"])
                        my_minutes = int((rain_ts - tick_ts) / 60)
                        my_was_active = my_minutes <= 60
                    else:
                        # Игрок вообще не писал до этого дождя
                        my_was_active = False
                except Exception:
                    pass

        merami = {
            "won": won,
            "pos": (recs.index(MY_NICK) + 1) if won else None,
            "minutes_ago": my_minutes,
            "was_active": my_was_active,
            "in_receivers": bool(my_rec),
        }
        out.append({
            "id": r["id"], "ts": r["ts"], "chat": r["chat"], "sender": r["sender"],
            "amount_text": r["amount_text"], "amount_rub": r.get("amount_rub") or 0,
            "amount_usd": r.get("amount_usd") or 0,
            "crypto_symbol": r.get("crypto_symbol") or "",
            "receivers_count": len(recs),
            "receivers": recs,
            "receivers_info": rinfo_sorted,
            "top20": top20,
            "merami": merami,
        })
    return {"hours": hours, "count": len(out), "rains": out}


@app.get("/api/chats_with_rains")
async def api_chats_with_rains(days: int = 90):
    since = time.time() - days * 86400
    rows = query_all("SELECT DISTINCT chat FROM events WHERE type='rain' AND ts >= ?", (since,))
    return {"chats": [r["chat"] for r in rows if r.get("chat")]}


@app.get("/api/server_info")
async def api_server_info():
    started = datetime.fromtimestamp(SERVER_START_TS, tz=timezone.utc)
    return {
        "started_at": SERVER_START_TS,
        "started_at_str": started.strftime("%Y-%m-%d %H:%M:%S UTC"),
        "uptime_sec": int(time.time() - SERVER_START_TS),
        "git_commit": SERVER_GIT_COMMIT,
        "git_branch": SERVER_GIT_BRANCH,
        "service_name": SERVER_SERVICE_NAME,
        "now_ts": time.time(),
    }


@app.get("/api/db_stats")
async def api_db_stats():
    tables = ["events", "bets", "intents", "bulk_alerts", "player_activity",
              "chat_players", "player_message_ticks", "moderator_messages",
              "moderators", "promocodes", "chat_tasks"]
    out = {}
    for t in tables:
        try:
            r = query_one(f"SELECT COUNT(*) AS c FROM {t}")
            out[t] = r["c"] if r else 0
        except Exception:
            out[t] = None
    try:
        rains = (query_one("SELECT COUNT(*) AS c FROM events WHERE type='rain'") or {}).get("c", 0)
        tips = (query_one("SELECT COUNT(*) AS c FROM events WHERE type='tip'") or {}).get("c", 0)
    except Exception:
        rains = tips = 0
    return {"tables": out, "rains": rains, "tips": tips, "now": time.time()}


@app.get("/api/daily_report")
async def api_daily_report(date: str):
    if not HAS_DOCX:
        raise HTTPException(503, "python-docx не установлен")
    try:
        dt = datetime.strptime(date, "%Y-%m-%d").replace(tzinfo=timezone.utc)
        start = dt.timestamp()
        end = start + 86400
    except Exception:
        raise HTTPException(400, "invalid date (YYYY-MM-DD)")

    events = query_all("""SELECT * FROM events WHERE ts>=? AND ts<=? ORDER BY ts ASC""",
                       (start, end))
    for r in events:
        try:
            r["receivers"] = json.loads(r.get("receivers") or "[]")
        except Exception:
            r["receivers"] = []

    rains = [e for e in events if e["type"] == "rain"]
    tips = [e for e in events if e["type"] == "tip"]
    total_rain_rub = sum((e.get("amount_rub") or 0) for e in rains)
    total_tip_rub = sum((e.get("amount_rub") or 0) for e in tips)

    send_rub = defaultdict(float)
    send_cnt = defaultdict(int)
    recv_rub = defaultdict(float)
    recv_cnt = defaultdict(int)
    my_share_rub = 0.0
    my_wins = 0
    for e in rains:
        s = e["sender"]
        send_rub[s] += (e.get("amount_rub") or 0)
        send_cnt[s] += 1
        recs = e.get("receivers") or []
        n = len(recs) or 1
        per = (e.get("amount_rub") or 0) / n
        for nick in recs:
            recv_rub[nick] += per
            recv_cnt[nick] += 1
        if MY_NICK in recs:
            my_share_rub += per
            my_wins += 1

    top_senders = sorted(send_rub.items(), key=lambda x: -x[1])[:15]
    top_receivers = sorted(recv_rub.items(), key=lambda x: -x[1])[:15]

    doc = Document()
    doc.add_heading(f"Shuffle — отчёт за {date}", level=0)
    doc.add_paragraph(f"Сгенерировано: {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M UTC')}")
    doc.add_heading("Сводка", level=1)
    doc.add_paragraph(f"Дождей: {len(rains)}")
    doc.add_paragraph(f"Раздали всего: {total_rain_rub:,.2f} ₽".replace(",", " "))
    doc.add_paragraph(f"Чаевых: {len(tips)}, сумма: {total_tip_rub:,.2f} ₽".replace(",", " "))
    doc.add_heading("Мой профиль", level=1)
    doc.add_paragraph(f"Ник: {MY_NICK}")
    doc.add_paragraph(f"Побед: {my_wins}")
    doc.add_paragraph(f"Получено: {my_share_rub:,.2f} ₽".replace(",", " "))

    doc.add_heading("Топ-15 раздающих", level=1)
    t = doc.add_table(rows=1, cols=3)
    t.style = "Light Grid Accent 1"
    hdr = t.rows[0].cells
    hdr[0].text = "Ник"; hdr[1].text = "Дождей"; hdr[2].text = "Сумма ₽"
    for nick, rub in top_senders:
        row = t.add_row().cells
        row[0].text = str(nick)
        row[1].text = str(send_cnt[nick])
        row[2].text = f"{rub:,.2f}".replace(",", " ")

    doc.add_heading("Топ-15 получателей", level=1)
    t2 = doc.add_table(rows=1, cols=3)
    t2.style = "Light Grid Accent 1"
    hdr2 = t2.rows[0].cells
    hdr2[0].text = "Ник"; hdr2[1].text = "Дождей"; hdr2[2].text = "Получено ₽"
    for nick, rub in top_receivers:
        row = t2.add_row().cells
        row[0].text = str(nick)
        row[1].text = str(recv_cnt[nick])
        row[2].text = f"{rub:,.2f}".replace(",", " ")

    doc.add_heading("Все дожди", level=1)
    for e in rains:
        tstr = datetime.fromtimestamp(e["ts"], tz=timezone.utc).strftime("%H:%M:%S")
        line = (f"{tstr} [{e['chat']}] {e['sender']} раздал "
                f"{e['amount_text']} {e.get('crypto_symbol') or ''} "
                f"на {len(e.get('receivers') or [])} чел")
        doc.add_paragraph(line)

    path = os.path.join(SCREENSHOT_DIR, f"report_{date}.docx")
    doc.save(path)
    return FileResponse(path,
                        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                        filename=f"shuffle_report_{date}.docx")


# ==================== API UI ====================
@app.get("/api/rates")
async def api_rates():
    with _server_rates_lock:
        return {"rates": dict(_server_rates), "updated_at": _server_rates_updated[0]}


@app.get("/api/status")
async def get_status():
    with _state_lock:
        instances = list(INSTANCE_STATES.values())
    instances.sort(key=lambda x: x.get("iid", 0))
    s = query_one("SELECT COUNT(*) AS c FROM events") or {"c": 0}
    tips = query_one("SELECT COUNT(*) AS c FROM events WHERE type='tip'") or {"c": 0}
    rains = query_one("SELECT COUNT(*) AS c FROM events WHERE type='rain'") or {"c": 0}
    first = query_one("SELECT MIN(ts) AS m FROM events WHERE ts IS NOT NULL") or {"m": None}
    bets_c = query_one("SELECT COUNT(*) AS c FROM bets") or {"c": 0}
    intents_c = query_one("SELECT COUNT(*) AS c FROM intents") or {"c": 0}
    return {"first_ts": first.get("m"), "total": s.get("c", 0),
            "tips": tips.get("c", 0), "rains": rains.get("c", 0),
            "bets": bets_c.get("c", 0), "intents": intents_c.get("c", 0),
            "instances": instances}


@app.post("/api/screenshot/{iid}")
async def request_screenshot(iid: int):
    with _state_lock:
        SCREENSHOT_REQUESTS[iid] = {"requested_at": time.time(), "done": False, "path": None}
    return {"ok": True}


@app.get("/api/screenshot/{iid}")
async def get_screenshot(iid: int):
    with _state_lock:
        r = SCREENSHOT_REQUESTS.get(iid)
    if not r or not r.get("done"):
        return JSONResponse(status_code=202, content={"pending": True})
    path = r.get("path")
    if not path or not os.path.exists(path):
        raise HTTPException(404, "not_available")
    return FileResponse(path, media_type="image/png")


@app.get("/api/events")
async def api_events(since: Optional[float] = None, until: Optional[float] = None,
                     type: Optional[str] = None, chat: Optional[str] = None,
                     sender: Optional[str] = None, min_rub: Optional[float] = None,
                     last_hours: Optional[int] = None, limit: int = 500):
    q = "SELECT * FROM events WHERE 1=1"
    p = []
    if last_hours:
        since = time.time() - last_hours * 3600
    if since is not None:
        q += " AND ts>=?"
        p.append(since)
    if until is not None:
        q += " AND ts<=?"
        p.append(until)
    if type:
        q += " AND type=?"
        p.append(type)
    if chat:
        q += " AND chat=?"
        p.append(chat)
    if sender:
        q += " AND sender LIKE ?"
        p.append(f"%{sender}%")
    if min_rub is not None:
        q += " AND amount_rub>=?"
        p.append(min_rub)
    q += (" ORDER BY ts DESC NULLS LAST, id DESC LIMIT ?" if USE_POSTGRES
          else " ORDER BY ts DESC, id DESC LIMIT ?")
    p.append(limit)
    rows = query_all(q, tuple(p))
    for r in rows:
        try:
            r["receivers"] = json.loads(r.get("receivers") or "[]")
        except Exception:
            r["receivers"] = []
        try:
            r["receivers_info"] = json.loads(r.get("receivers_info") or "[]")
        except Exception:
            r["receivers_info"] = []
        rc = len(r["receivers"]) or 1
        r["per_person_rub"] = (r.get("amount_rub") or 0) / rc
    return {"events": rows}


@app.get("/api/receivers_top")
async def api_receivers_top(limit: int = 100, metric: str = "total", days: int = 30,
                            window_min: int = 25):
    since = time.time() - days * 86400
    rows = query_all("""SELECT chat,ts,amount_rub,receivers,receivers_info FROM events
                        WHERE type='rain' AND ts>=? ORDER BY ts DESC""", (since,))
    stats = defaultdict(lambda: {"total_rub": 0.0, "wins": 0, "active_wins": 0,
                                  "unknown_activity_wins": 0,
                                  "chats": defaultdict(int), "last_seen": 0.0})
    for r in rows:
        try:
            recs = json.loads(r.get("receivers") or "[]")
        except Exception:
            recs = []
        if not recs:
            continue
        try:
            rinfo = json.loads(r.get("receivers_info") or "[]")
        except Exception:
            rinfo = []
        rmap = {x.get("name"): x for x in rinfo if isinstance(x, dict) and x.get("name")}
        amt = r.get("amount_rub") or 0
        per = amt / len(recs)
        chat = r.get("chat") or "?"
        for nick in recs:
            if not nick:
                continue
            s = stats[nick]
            s["total_rub"] += per
            s["wins"] += 1
            s["chats"][chat] += 1
            info = rmap.get(nick)
            if info:
                if info.get("wrote_in_25min"):
                    s["active_wins"] += 1
                try:
                    ls = float(info.get("last_msg_ts") or 0)
                except Exception:
                    ls = 0
                if ls > s["last_seen"]:
                    s["last_seen"] = ls
            else:
                s["unknown_activity_wins"] += 1
    arr = []
    for nick, s in stats.items():
        known = s["wins"] - s["unknown_activity_wins"]
        pct = (s["active_wins"] / known * 100) if known > 0 else None
        arr.append({"nickname": nick, "total_rub": round(s["total_rub"], 2),
                    "wins": s["wins"], "active_wins": s["active_wins"],
                    "known_wins": known,
                    "active_percent": round(pct, 1) if pct is not None else None,
                    "chats": dict(s["chats"]), "last_seen": s["last_seen"]})
    if metric == "count":
        arr.sort(key=lambda x: x["wins"], reverse=True)
    elif metric == "active_percent":
        arr.sort(key=lambda x: (x["active_percent"] if x["active_percent"] is not None else -1,
                                 x["wins"]), reverse=True)
    else:
        arr.sort(key=lambda x: x["total_rub"], reverse=True)
    return {"players": arr[:limit], "total": len(arr),
            "window_min": window_min, "days": days}


@app.get("/api/chart")
async def api_chart(type: Optional[str] = None, chat: Optional[str] = None,
                    since: Optional[float] = None, until: Optional[float] = None,
                    last_hours: Optional[int] = None):
    if last_hours:
        since = time.time() - last_hours * 3600
    hour_expr = ("CAST(EXTRACT(HOUR FROM to_timestamp(ts)) AS INTEGER)" if USE_POSTGRES
                 else "CAST(strftime('%H', ts, 'unixepoch') AS INTEGER)")
    q = f"SELECT {hour_expr} AS h, COUNT(*) AS c, COALESCE(SUM(amount_rub),0) AS s FROM events WHERE ts IS NOT NULL"
    p = []
    if since is not None:
        q += " AND ts>=?"
        p.append(since)
    if until is not None:
        q += " AND ts<=?"
        p.append(until)
    if type:
        q += " AND type=?"
        p.append(type)
    if chat:
        q += " AND chat=?"
        p.append(chat)
    q += " GROUP BY h"
    rows = query_all(q, tuple(p))
    buckets = {h: {"count": 0, "sum": 0.0} for h in range(24)}
    for r in rows:
        if r.get("h") is not None:
            buckets[int(r["h"])] = {"count": r["c"], "sum": r["s"] or 0}
    return {"buckets": buckets}


@app.get("/api/days")
async def api_days(days: int = 90):
    since = time.time() - days * 86400
    day_expr = ("to_char(to_timestamp(ts),'YYYY-MM-DD')" if USE_POSTGRES
                else "strftime('%Y-%m-%d',ts,'unixepoch')")
    rows = query_all(f"""SELECT {day_expr} AS day, COUNT(*) AS cnt,
                         SUM(CASE WHEN type='rain' THEN 1 ELSE 0 END) AS rains,
                         SUM(CASE WHEN type='tip' THEN 1 ELSE 0 END) AS tips,
                         COALESCE(SUM(amount_rub),0) AS total_rub
                         FROM events WHERE ts IS NOT NULL AND ts>=?
                         GROUP BY day ORDER BY day DESC""", (since,))
    return {"days": rows}


@app.get("/api/day_detail")
async def api_day_detail(date: str):
    try:
        dt = datetime.strptime(date, "%Y-%m-%d").replace(tzinfo=timezone.utc)
        start = dt.timestamp()
        end = start + 86400
    except Exception:
        raise HTTPException(400, "invalid date")
    events = query_all("""SELECT * FROM events WHERE ts>=? AND ts<=? ORDER BY ts ASC""",
                       (start, end))
    for r in events:
        try:
            r["receivers"] = json.loads(r.get("receivers") or "[]")
        except Exception:
            r["receivers"] = []
    recv = defaultdict(lambda: {"count": 0, "rub": 0.0, "usd": 0.0})
    send = defaultdict(lambda: {"count": 0, "rub": 0.0, "usd": 0.0})
    for e in events:
        if e["type"] == "rain":
            s = send[e["sender"]]
            s["count"] += 1
            s["rub"] += (e.get("amount_rub") or 0)
            s["usd"] += (e.get("amount_usd") or 0)
            recs = e.get("receivers") or []
            n = len(recs) or 1
            for nick in recs:
                r = recv[nick]
                r["count"] += 1
                r["rub"] += (e.get("amount_rub") or 0) / n
                r["usd"] += (e.get("amount_usd") or 0) / n
    return {"date": date, "events_count": len(events), "events": events[:500],
            "top_receivers": sorted([{"nickname": k, **v} for k, v in recv.items()],
                                     key=lambda x: -x["rub"])[:30],
            "top_senders": sorted([{"nickname": k, **v} for k, v in send.items()],
                                   key=lambda x: -x["rub"])[:30]}


@app.get("/api/chats")
async def api_chats():
    rows = query_all("SELECT DISTINCT chat FROM events WHERE chat IS NOT NULL ORDER BY chat")
    existing = [r["chat"] for r in rows]
    for c in ALL_CHATS:
        if c not in existing:
            existing.append(c)
    return {"chats": sorted(existing)}


@app.get("/api/contexts")
async def api_contexts(chat: Optional[str] = None, limit: int = 30):
    q = ("SELECT id,ts,created_at,chat,sender,receivers,receivers_info,amount_text,"
         "amount_rub,amount_usd,crypto_symbol,context_messages,context_bets "
         "FROM events WHERE type='rain'")
    p = []
    if chat:
        q += " AND chat=?"
        p.append(chat)
    q += (" ORDER BY ts DESC NULLS LAST, id DESC LIMIT ?" if USE_POSTGRES
          else " ORDER BY ts DESC, id DESC LIMIT ?")
    p.append(limit)
    rows = query_all(q, tuple(p))
    for r in rows:
        for k in ("receivers", "receivers_info", "context_messages", "context_bets"):
            try:
                r[k] = json.loads(r.get(k) or "[]")
            except Exception:
                r[k] = []
    return {"contexts": rows}


@app.get("/api/intents")
async def api_intents(chat: Optional[str] = None, sender: Optional[str] = None,
                      only_matched: Optional[int] = None, limit: int = 100):
    q = "SELECT * FROM intents WHERE 1=1"
    p = []
    if chat:
        q += " AND chat=?"
        p.append(chat)
    if sender:
        q += " AND sender LIKE ?"
        p.append(f"%{sender}%")
    if only_matched == 1:
        q += " AND matched_rain_id IS NOT NULL"
    if only_matched == 0:
        q += " AND matched_rain_id IS NULL"
    q += (" ORDER BY ts DESC NULLS LAST, id DESC LIMIT ?" if USE_POSTGRES
          else " ORDER BY ts DESC, id DESC LIMIT ?")
    p.append(limit)
    return {"intents": query_all(q, tuple(p))}


@app.get("/api/intents_stats")
async def api_intents_stats():
    if USE_POSTGRES:
        rows = query_all("""SELECT sender AS nickname, COUNT(*) AS total,
                            SUM(CASE WHEN matched_rain_id IS NOT NULL THEN 1 ELSE 0 END) AS matched,
                            ROUND(AVG(percent),0) AS avg_percent,
                            ROUND(CAST(SUM(CASE WHEN matched_rain_id IS NOT NULL THEN 1 ELSE 0 END) AS NUMERIC)
                                  /CAST(COUNT(*) AS NUMERIC)*100,1) AS percent
                            FROM intents GROUP BY sender HAVING COUNT(*)>=2
                            ORDER BY percent DESC, matched DESC LIMIT 50""")
    else:
        rows = query_all("""SELECT sender AS nickname, COUNT(*) AS total,
                            SUM(CASE WHEN matched_rain_id IS NOT NULL THEN 1 ELSE 0 END) AS matched,
                            ROUND(AVG(percent),0) AS avg_percent,
                            ROUND(CAST(SUM(CASE WHEN matched_rain_id IS NOT NULL THEN 1 ELSE 0 END) AS REAL)
                                  /CAST(COUNT(*) AS REAL)*100,1) AS percent
                            FROM intents GROUP BY sender HAVING COUNT(*)>=2
                            ORDER BY percent DESC, matched DESC LIMIT 50""")
    return {"stats": rows}


@app.get("/api/bulk")
async def api_bulk(since: Optional[float] = None, until: Optional[float] = None,
                   senders: Optional[str] = None, limit: int = 500):
    q = "SELECT * FROM bulk_alerts WHERE 1=1"
    p = []
    if since is not None:
        q += " AND ts>=?"
        p.append(since)
    if until is not None:
        q += " AND ts<=?"
        p.append(until)
    if senders:
        names = [s.strip() for s in senders.split(",") if s.strip()]
        if names:
            phs = ",".join(["?"] * len(names))
            q += f" AND sender IN ({phs})"
            p.extend(names)
    q += (" ORDER BY ts DESC NULLS LAST, id DESC LIMIT ?" if USE_POSTGRES
          else " ORDER BY ts DESC, id DESC LIMIT ?")
    p.append(limit)
    rows = query_all(q, tuple(p))
    for r in rows:
        try:
            r["receivers"] = json.loads(r.get("receivers") or "[]")
        except Exception:
            r["receivers"] = []
    return {"alerts": rows}


@app.get("/api/player")
async def api_player(username: Optional[str] = None, chat: Optional[str] = None,
                     since: Optional[float] = None, limit: int = 500):
    q = "SELECT * FROM player_activity WHERE 1=1"
    p = []
    if username:
        q += " AND username LIKE ?"
        p.append(f"%{username}%")
    if chat:
        q += " AND chat=?"
        p.append(chat)
    if since is not None:
        q += " AND ts>=?"
        p.append(since)
    q += (" ORDER BY ts DESC NULLS LAST, id DESC LIMIT ?" if USE_POSTGRES
          else " ORDER BY ts DESC, id DESC LIMIT ?")
    p.append(limit)
    return {"rows": query_all(q, tuple(p)), "hourly": {h: 0 for h in range(24)}}


@app.get("/api/players")
async def api_players(chat: Optional[str] = None, username: Optional[str] = None,
                      online_only: Optional[int] = None, min_msgs: Optional[int] = None,
                      limit: int = 500):
    q = "SELECT chat,username,msg_count,first_seen,last_seen FROM chat_players WHERE 1=1"
    p = []
    if chat:
        q += " AND chat=?"
        p.append(chat)
    if username:
        q += " AND username LIKE ?"
        p.append(f"%{username}%")
    if min_msgs is not None:
        q += " AND msg_count>=?"
        p.append(min_msgs)
    if online_only == 1:
        q += " AND last_seen>=?"
        p.append(time.time() - 300)
    q += (" ORDER BY last_seen DESC NULLS LAST, msg_count DESC LIMIT ?" if USE_POSTGRES
          else " ORDER BY last_seen DESC, msg_count DESC LIMIT ?")
    p.append(limit)
    return {"players": query_all(q, tuple(p))}


@app.post("/api/players/relocate")
async def api_players_relocate(body: dict):
    username = (body.get("username") or "").strip()
    new_chat = (body.get("chat") or "").strip()
    if not username or not new_chat:
        raise HTTPException(400, "username, chat required")
    p = ph()
    if USE_POSTGRES:
        execute(f"""INSERT INTO chat_players (chat,username,msg_count,first_seen,last_seen)
                    VALUES ({p},{p},0,{p},{p}) ON CONFLICT (chat,username) DO NOTHING""",
                (new_chat, username, time.time(), time.time()))
    else:
        execute(f"""INSERT OR IGNORE INTO chat_players
                    (chat,username,msg_count,first_seen,last_seen)
                    VALUES (?,?,0,?,?)""",
                (new_chat, username, time.time(), time.time()))
    return {"ok": True}


@app.post("/api/players/set_chats")
async def api_players_set_chats(body: dict):
    username = (body.get("username") or "").strip()
    chats = body.get("chats") or []
    if not username:
        raise HTTPException(400, "username required")
    chats = [c.strip().upper() for c in chats if c.strip()]
    execute("DELETE FROM chat_players WHERE username=?", (username,))
    p = ph()
    now = time.time()
    for c in chats:
        if USE_POSTGRES:
            execute(f"""INSERT INTO chat_players (chat,username,msg_count,first_seen,last_seen)
                        VALUES ({p},{p},0,{p},{p}) ON CONFLICT (chat,username) DO NOTHING""",
                    (c, username, now, now))
        else:
            execute(f"""INSERT OR IGNORE INTO chat_players
                        (chat,username,msg_count,first_seen,last_seen)
                        VALUES (?,?,0,?,?)""",
                    (c, username, now, now))
    return {"ok": True, "chats": chats}


@app.get("/api/players_by_chat")
async def api_players_by_chat():
    cutoff = time.time() - 300
    rows = query_all("""SELECT chat, COUNT(*) AS total_players,
                        SUM(CASE WHEN last_seen>=? THEN 1 ELSE 0 END) AS online_players,
                        SUM(msg_count) AS total_messages
                        FROM chat_players GROUP BY chat ORDER BY total_players DESC""",
                     (cutoff,))
    return {"chats": rows}


@app.get("/api/players_by_language")
async def api_players_by_language():
    if USE_POSTGRES:
        rows = query_all("""SELECT username, STRING_AGG(chat,',') AS chats_raw,
                            COUNT(DISTINCT chat) AS chat_count,
                            SUM(msg_count) AS total_msgs, MAX(last_seen) AS last_seen
                            FROM chat_players GROUP BY username
                            ORDER BY SUM(msg_count) DESC LIMIT 5000""")
    else:
        rows = query_all("""SELECT username, GROUP_CONCAT(chat) AS chats_raw,
                            COUNT(DISTINCT chat) AS chat_count,
                            SUM(msg_count) AS total_msgs, MAX(last_seen) AS last_seen
                            FROM chat_players GROUP BY username
                            ORDER BY SUM(msg_count) DESC LIMIT 5000""")
    out = []
    for r in rows:
        chats = list(set((r.get("chats_raw") or "").split(",")))
        r["chats"] = [c.strip() for c in chats if c.strip()]
        r.pop("chats_raw", None)
        out.append(r)
    return {"players": out, "all_chats": ALL_CHATS}


@app.get("/api/bets_to_rains")
async def api_bets_to_rains(window_sec: int = BET_TO_RAIN_WINDOW_SEC,
                             last_hours: Optional[int] = None):
    since = time.time() - (last_hours * 3600 if last_hours else 7 * 86400)
    cleared = float(meta_get("bets_to_rains_cleared_at") or 0)
    since = max(since, cleared)

    bets = query_all("""SELECT id,ts,sender FROM bets
                        WHERE ts>=? AND COALESCE(amount_usd,0)>=20
                        ORDER BY ts ASC""", (since,))
    rains = query_all("""SELECT ts FROM events WHERE type='rain' AND ts>=? AND ts IS NOT NULL
                        ORDER BY ts ASC""", (since,))
    rlist = sorted([r["ts"] for r in rains if r.get("ts")])
    stats = defaultdict(lambda: {"bets": 0, "triggered": 0, "total_rain_rub": 0.0})
    import bisect
    for b in bets:
        s = (b.get("sender") or "").strip()
        if not s:
            continue
        stats[s]["bets"] += 1
        t0 = b["ts"]
        t1 = t0 + window_sec
        idx = bisect.bisect_left(rlist, t0)
        if idx < len(rlist) and rlist[idx] <= t1:
            stats[s]["triggered"] += 1
    arr = []
    for nick, d in stats.items():
        if d["bets"] < 3:
            continue
        pct = (d["triggered"] / d["bets"] * 100) if d["bets"] else 0
        arr.append({"nickname": nick, "bets": d["bets"],
                    "triggered": d["triggered"], "percent": round(pct, 1)})
    arr.sort(key=lambda x: (-x["percent"], -x["bets"]))
    return {"stats": arr[:200], "window_sec": window_sec, "cleared_at": cleared}


@app.get("/api/bets")
async def api_bets(since: Optional[float] = None, until: Optional[float] = None,
                   chat: Optional[str] = None, sender: Optional[str] = None,
                   min_usd: Optional[float] = None, source: Optional[str] = None,
                   last_hours: Optional[int] = None, limit: int = 500):
    if last_hours:
        since = time.time() - last_hours * 3600
    q = "SELECT * FROM bets WHERE 1=1"
    p = []
    if since is not None:
        q += " AND ts>=?"
        p.append(since)
    if until is not None:
        q += " AND ts<=?"
        p.append(until)
    if chat:
        q += " AND chat=?"
        p.append(chat)
    if sender:
        q += " AND (sender LIKE ? OR real_sender LIKE ?)"
        p.extend([f"%{sender}%", f"%{sender}%"])
    if min_usd is not None:
        q += " AND amount_usd>=?"
        p.append(min_usd)
    if source:
        q += " AND source=?"
        p.append(source)
    q += (" ORDER BY ts DESC NULLS LAST, id DESC LIMIT ?" if USE_POSTGRES
          else " ORDER BY ts DESC, id DESC LIMIT ?")
    p.append(limit)
    return {"bets": query_all(q, tuple(p))}


@app.get("/api/bets_stats")
async def api_bets_stats():
    since = time.time() - 86400
    rows = query_all("""SELECT chat, COUNT(*) AS cnt,
                        COALESCE(SUM(amount_usd),0) AS total_usd,
                        COALESCE(AVG(amount_usd),0) AS avg_usd
                        FROM bets WHERE ts>=? GROUP BY chat ORDER BY total_usd DESC""",
                     (since,))
    return {"stats": rows}


@app.get("/api/bets_top_senders")
async def api_bets_top_senders(last_hours: Optional[int] = None):
    cleared = float(meta_get("bets_top_cleared_at") or 0)
    since = time.time() - (last_hours * 3600 if last_hours else 7 * 86400)
    since = max(since, cleared)
    rows = query_all("""SELECT sender AS nickname, COUNT(*) AS cnt,
                        COALESCE(SUM(amount_usd),0) AS total_usd
                        FROM bets WHERE ts>=? GROUP BY sender
                        ORDER BY total_usd DESC LIMIT 30""", (since,))
    return {"senders": rows, "cleared_at": cleared, "since": since}


@app.get("/api/health")
async def health():
    return {"ok": True, "groq": bool(GROQ_API_KEY and GROQ_API_KEY.startswith("gsk_"))}


@app.get("/api/stats")
async def get_stats():
    total_rains = (query_one("SELECT COUNT(*) AS c FROM events WHERE type='rain'") or {}).get("c", 0)
    total_tips = (query_one("SELECT COUNT(*) AS c FROM events WHERE type='tip'") or {}).get("c", 0)
    trub = (query_one("SELECT COALESCE(SUM(amount_rub),0) AS s FROM events WHERE type='rain'") or {}).get("s", 0)
    ttrub = (query_one("SELECT COALESCE(SUM(amount_rub),0) AS s FROM events WHERE type='tip'") or {}).get("s", 0)
    recent = query_all("""SELECT type,chat,sender,receivers,amount_text,crypto_symbol,
                          amount_rub,amount_usd,created_at,ts
                          FROM events ORDER BY id DESC LIMIT 30""")
    for r in recent:
        try:
            rcs = json.loads(r.get("receivers") or "[]")
        except Exception:
            rcs = []
        r["receivers"] = rcs
        r["per_person_rub"] = (r.get("amount_rub") or 0) / max(len(rcs), 1)
    return {"totals": {"rains": total_rains, "tips": total_tips,
                       "rain_rub": round(trub or 0, 2),
                       "tip_rub": round(ttrub or 0, 2)},
            "recent": recent,
            "generated_at": datetime.now(timezone.utc).isoformat()}


@app.post("/api/clear")
async def clear_database(body: dict, x_api_key: str = Header(default="")):
    if x_api_key != API_KEY:
        raise HTTPException(401, "Invalid API key")
    if body.get("confirm") != "DELETE_ALL":
        raise HTTPException(400, "Подтверждение обязательно")
    if USE_POSTGRES:
        execute("TRUNCATE TABLE events RESTART IDENTITY CASCADE;")
    else:
        execute("DELETE FROM events;")
    return {"ok": True, "message": "Все данные удалены"}


# ==================== HTML ====================
BASE_HTML = """<!DOCTYPE html>
<html lang="ru"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>{title}</title>
<style>
body{{margin:0;background:#0e1116;color:#e6e9ef;font-family:-apple-system,Segoe UI,Roboto,Arial,sans-serif}}
a{{color:#7cc4ff;text-decoration:none}}
header{{background:#161b22;border-bottom:1px solid #26303a;padding:14px 24px;display:flex;align-items:center;gap:24px;flex-wrap:wrap}}
.brand{{font-weight:700;font-size:18px;color:#fff}}
nav a{{margin-right:16px;color:#c9d1d9}}
nav a.active{{color:#7cc4ff;font-weight:600}}
main{{max-width:1200px;margin:0 auto;padding:24px}}
.grid{{display:grid;grid-template-columns:repeat(auto-fit,minmax(220px,1fr));gap:16px}}
.card{{background:#161b22;border:1px solid #26303a;border-radius:12px;padding:16px 20px}}
.card h3{{margin:0 0 8px;color:#8b949e;font-weight:500;font-size:13px;text-transform:uppercase;display:flex;align-items:center;gap:6px}}
.card .big{{font-size:26px;font-weight:700;color:#fff}}
table{{width:100%;border-collapse:collapse}}
th,td{{padding:10px 12px;text-align:left;border-bottom:1px solid #26303a}}
th{{color:#8b949e;font-weight:500;font-size:13px}}
td.amount{{color:#4ade80;font-weight:600}}
tr:hover td{{background:#1c222b}}
.tag{{display:inline-block;background:#21262d;padding:2px 8px;border-radius:6px;font-size:11px;color:#8b949e;margin-left:6px}}
.muted{{color:#8b949e}}
.chat-pill{{display:inline-block;background:#21262d;padding:2px 8px;border-radius:6px;font-size:11px;color:#7cc4ff}}
.hint-btn{{display:inline-block;width:18px;height:18px;line-height:18px;text-align:center;border-radius:50%;
  background:#1e88e5;color:#fff;font-size:12px;cursor:help;font-weight:bold;
  position:relative;vertical-align:middle;flex-shrink:0;user-select:none}}
.hint-btn:hover{{background:#42a5f5}}
.hint-btn:hover::after{{content:attr(data-hint);position:absolute;top:26px;left:-200px;width:420px;
  background:#1a1f28;color:#e6e6e6;border:1px solid #2f3d55;border-radius:8px;padding:14px;
  font-size:12px;font-weight:normal;line-height:1.6;z-index:1000;white-space:pre-wrap;
  box-shadow:0 8px 30px rgba(0,0,0,.8);text-align:left;pointer-events:none}}
</style></head><body>
<header>
  <div class="brand">🌧 SHUFFLE RAIN <small>статистика</small></div>
  <nav>
    <a href="/" class="{active_home}">Главная</a>
    <a href="/hosts" class="{active_hosts}">Раздающие</a>
    <a href="/receivers" class="{active_receivers}">Получатели</a>
    <a href="/monitor">Мониторинг</a>
  </nav>
</header>
<main>{content}</main>
</body></html>"""


def fmt_rub(x):
    try:
        return f"{x:,.2f}".replace(",", " ").replace(".", ",") + " ₽"
    except Exception:
        return "0 ₽"


@app.get("/", response_class=HTMLResponse)
async def page_home():
    t = (query_one("SELECT COUNT(*) AS c FROM events WHERE type='rain'") or {}).get("c", 0)
    tips = (query_one("SELECT COUNT(*) AS c FROM events WHERE type='tip'") or {}).get("c", 0)
    tr = (query_one("SELECT COALESCE(SUM(amount_rub),0) AS s FROM events WHERE type='rain'") or {}).get("s", 0)
    ttr = (query_one("SELECT COALESCE(SUM(amount_rub),0) AS s FROM events WHERE type='tip'") or {}).get("s", 0)
    content = f"""
    <div class="grid">
      <div class="card"><h3>Дождей</h3><div class="big">{t}</div></div>
      <div class="card"><h3>Раздали всего</h3><div class="big">{fmt_rub(tr or 0)}</div></div>
      <div class="card"><h3>Чаевых</h3><div class="big">{tips}</div></div>
      <div class="card"><h3>Чаевых сумма</h3><div class="big">{fmt_rub(ttr or 0)}</div></div>
    </div>
    <div class="grid" style="margin-top:16px">
      <div class="card"><h3>Последний деплой сервера</h3><div id="deployBlock" class="big" style="font-size:16px">—</div></div>
    </div>
    <div class="card" style="margin-top:16px"><h3>Последние события</h3><div id="recent"></div></div>
    <script>
      function esc(s){{return (s==null?'':String(s)).replace(/[<>&"]/g,c=>({{'<':'&lt;','>':'&gt;','&':'&amp;','"':'&quot;'}}[c]));}}
      function fmtR(v){{return (v||0).toFixed(0);}}
      async function loadOnce() {{ const r = await fetch('/api/stats'); const d = await r.json(); renderRecent(d.recent); }}
      async function loadDeployInfo() {{
        try {{
          const r = await fetch('/api/server_info').then(r=>r.json());
          const started = new Date(r.started_at * 1000);
          const upMin = Math.round(r.uptime_sec / 60);
          const ago = upMin < 60 ? upMin + 'м' : Math.floor(upMin/60) + 'ч ' + (upMin % 60) + 'м';
          document.getElementById('deployBlock').innerHTML =
            started.toLocaleString('ru-RU') + ' <span class="muted">(' + ago + ' назад)</span>'
            + '<div class="muted" style="font-size:12px;margin-top:4px">commit: <b>' + (r.git_commit||'?') + '</b> · branch: ' + (r.git_branch||'-') + '</div>';
        }} catch(e) {{}}
      }}
      function renderRecent(rows) {{
        const el = document.getElementById('recent'); if (!rows) return;
        el.innerHTML = rows.map(r => {{
          const tag = r.type === 'rain' ? 'RAIN' : 'TIP';
          const amt = r.amount_text || '';
          const rub = r.amount_rub ? ' (' + fmtR(r.amount_rub) + ' ₽)' : '';
          const per = r.per_person_rub ? ' <span style="color:#8ab4f8;font-size:11px">по ' + fmtR(r.per_person_rub) + ' ₽</span>' : '';
          const recs = (r.receivers || []);
          const recsHtml = recs.slice(0,5).map(esc).join(', ') + (recs.length > 5 ? ` ...и ещё ${{recs.length-5}}` : '');
          return `<div style="padding:8px 0;border-bottom:1px solid #26303a;font-size:13px"><span class="tag">${{tag}}</span> <span class="chat-pill">${{r.chat}}</span> <b style="margin-left:6px">${{esc(r.sender)}}</b><span class="muted"> → ${{recsHtml}}</span>${{per}}<span style="float:right;color:#4ade80">${{esc(amt)}}${{rub}}</span></div>`;
        }}).join('');
      }}
      loadOnce(); setInterval(loadOnce, 15000);
      loadDeployInfo(); setInterval(loadDeployInfo, 30000);
    </script>"""
    return BASE_HTML.format(title="Shuffle Rain Stats", active_home="active",
                            active_hosts="", active_receivers="", content=content)


@app.get("/receivers", response_class=HTMLResponse)
async def page_receivers():
    content = """
    <div class="card">
      <h3 style="display:flex;align-items:center;gap:6px">🎯 Топ активных получателей
        <span class="hint-btn" data-hint="Кто чаще всего получает дожди.&#10;Активный = писал в том же чате за 25 мин до дождя.">?</span></h3>
      <div style="display:flex;gap:8px;flex-wrap:wrap;margin-bottom:12px">
        <select id="rtMetric" onchange="loadReceiversTop()"><option value="total">По сумме ₽</option><option value="count">По кол-ву</option><option value="active_percent">По % активных</option></select>
        <select id="rtLimit" onchange="loadReceiversTop()"><option value="20">Топ-20</option><option value="50">Топ-50</option><option value="100" selected>Топ-100</option><option value="200">Топ-200</option><option value="500">Все</option></select>
        <select id="rtDays" onchange="loadReceiversTop()"><option value="7">7 дней</option><option value="30" selected>30 дней</option><option value="90">90 дней</option></select>
        <button onclick="loadReceiversTop()">Обновить</button>
        <span id="rtStatus" class="muted"></span>
      </div>
      <table id="rtTable"><thead><tr>
        <th onclick="sortSimple('rtTable',0,1)" style="cursor:pointer">#</th>
        <th onclick="sortSimple('rtTable',1,1)" style="cursor:pointer">Ник</th>
        <th onclick="sortSimple('rtTable',2,1)" style="cursor:pointer">Получил ₽</th>
        <th onclick="sortSimple('rtTable',3,1)" style="cursor:pointer">Дождей</th>
        <th onclick="sortSimple('rtTable',4,1)" style="cursor:pointer">Активных</th>
        <th onclick="sortSimple('rtTable',5,1)" style="cursor:pointer">%</th>
        <th>Чаты</th>
      </tr></thead><tbody id="rtBody"><tr><td colspan="7" class="muted">Загрузка...</td></tr></tbody></table>
    </div>
    <script>
      const esc = s => (s==null?'':String(s)).replace(/[<>&"]/g, c=>({'<':'&lt;','>':'&gt;','&':'&amp;','"':'&quot;'}[c]));
      const fmtMoney = v => (v||0).toLocaleString('ru-RU',{maximumFractionDigits:2});
      async function loadReceiversTop() {
        const m = document.getElementById('rtMetric').value;
        const l = document.getElementById('rtLimit').value;
        const d = document.getElementById('rtDays').value;
        document.getElementById('rtStatus').textContent = '⏳...';
        const r = await fetch('/api/receivers_top?metric='+m+'&limit='+l+'&days='+d).then(r=>r.json());
        const body = document.getElementById('rtBody');
        if (!r.players.length) { body.innerHTML = '<tr><td colspan="7" class="muted">Нет данных</td></tr>'; document.getElementById('rtStatus').textContent=''; return; }
        body.innerHTML = r.players.map((p,i) => {
          const chats = Object.entries(p.chats||{}).sort((a,b)=>b[1]-a[1]).slice(0,6).map(([c,n])=>'<span class="tag">'+esc(c)+':'+n+'</span>').join(' ');
          let pct = '<span class="muted">—</span>';
          if (p.active_percent !== null && p.active_percent !== undefined) {
            const col = p.active_percent >= 60 ? '#4ade80' : (p.active_percent >= 30 ? '#fbbf24' : '#f87171');
            pct = '<span style="color:'+col+';font-weight:600">'+p.active_percent+'%</span>';
          }
          return `<tr><td data-v="${i+1}">${i+1}</td><td data-v="${esc(p.nickname)}"><b>${esc(p.nickname)}</b></td><td data-v="${p.total_rub}" class="amount">${fmtMoney(p.total_rub)} ₽</td><td data-v="${p.wins}">${p.wins}</td><td data-v="${p.active_wins}">${p.active_wins}/${p.known_wins}</td><td data-v="${p.active_percent||-1}">${pct}</td><td>${chats}</td></tr>`;
        }).join('');
        document.getElementById('rtStatus').textContent = 'Найдено: '+r.total;
      }
      function sortSimple(tableId, colIdx, dir) {
        const tb = document.getElementById(tableId).querySelector('tbody');
        const rows = Array.from(tb.querySelectorAll('tr'));
        rows.sort((a,b)=>{
          const av = a.cells[colIdx]?.getAttribute('data-v') || a.cells[colIdx]?.textContent.trim() || '';
          const bv = b.cells[colIdx]?.getAttribute('data-v') || b.cells[colIdx]?.textContent.trim() || '';
          const an = parseFloat(av), bn = parseFloat(bv);
          if (!isNaN(an) && !isNaN(bn)) return dir * (an - bn);
          return dir * av.localeCompare(bv);
        });
        rows.forEach(r => tb.appendChild(r));
      }
      loadReceiversTop();
    </script>"""
    return BASE_HTML.format(title="Shuffle — Получатели", active_home="",
                            active_hosts="", active_receivers="active", content=content)


@app.get("/hosts", response_class=HTMLResponse)
async def page_hosts():
    r1 = query_all("""SELECT sender AS nickname, COUNT(*) AS cnt,
                      COALESCE(SUM(amount_rub),0) AS total
                      FROM events WHERE type='rain' GROUP BY sender
                      ORDER BY total DESC LIMIT 50""")
    r2 = query_all("""SELECT sender AS nickname, COUNT(*) AS cnt,
                      COALESCE(SUM(amount_rub),0) AS total
                      FROM events WHERE type='rain' GROUP BY sender
                      ORDER BY cnt DESC LIMIT 50""")
    t1 = query_all("""SELECT sender AS nickname, COUNT(*) AS cnt,
                      COALESCE(SUM(amount_rub),0) AS total
                      FROM events WHERE type='tip' GROUP BY sender
                      ORDER BY total DESC LIMIT 50""")
    t2 = query_all("""SELECT sender AS nickname, COUNT(*) AS cnt,
                      COALESCE(SUM(amount_rub),0) AS total
                      FROM events WHERE type='tip' GROUP BY sender
                      ORDER BY cnt DESC LIMIT 50""")

    def table(rows):
        if not rows:
            return "<tr><td colspan='4' class='muted'>Пусто</td></tr>"
        html = ""
        for i, r in enumerate(rows, 1):
            html += f"""<tr><td>{i}</td><td><b>{r['nickname']}</b></td><td>{r['cnt']}</td><td class="amount">{fmt_rub(r['total'])}</td></tr>"""
        return html

    content = f"""
    <div class="grid" style="grid-template-columns: 1fr 1fr">
      <div class="card"><h3>🌧 Раздал дождей — по сумме</h3><table><thead><tr><th>#</th><th>Ник</th><th>Кол-во</th><th>Сумма</th></tr></thead><tbody>{table(r1)}</tbody></table></div>
      <div class="card"><h3>🌧 Раздал дождей — по кол-ву</h3><table><thead><tr><th>#</th><th>Ник</th><th>Кол-во</th><th>Сумма</th></tr></thead><tbody>{table(r2)}</tbody></table></div>
      <div class="card"><h3>💸 Отправил чаевых — по сумме</h3><table><thead><tr><th>#</th><th>Ник</th><th>Кол-во</th><th>Сумма</th></tr></thead><tbody>{table(t1)}</tbody></table></div>
      <div class="card"><h3>💸 Отправил чаевых — по кол-ву</h3><table><thead><tr><th>#</th><th>Ник</th><th>Кол-во</th><th>Сумма</th></tr></thead><tbody>{table(t2)}</tbody></table></div>
    </div>"""
    return BASE_HTML.format(title="Shuffle — Раздающие", active_home="",
                            active_hosts="active", active_receivers="", content=content)


# ==================== MONITOR ====================
MONITOR_HTML = r"""<!DOCTYPE html>
<html lang="ru"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>Shuffle Monitor</title>
<script src="https://cdn.jsdelivr.net/npm/chart.js@4.4.1/dist/chart.umd.min.js"></script>
<style>
*{box-sizing:border-box}
body{margin:0;background:#0f1115;color:#e6e6e6;font:14px/1.4 -apple-system,Segoe UI,Roboto,sans-serif}
header{display:flex;align-items:center;gap:16px;padding:12px 20px;background:#171a21;border-bottom:1px solid #2a2f3a;position:sticky;top:0;z-index:10;flex-wrap:wrap}
header h1{font-size:16px;margin:0;font-weight:600}
header a{color:#7cc4ff;text-decoration:none;font-size:13px}
.tabs{display:flex;gap:4px;padding:10px 20px;border-bottom:1px solid #2a2f3a;background:#12151b;flex-wrap:wrap}
.tabs button{background:transparent;color:#aab;border:1px solid transparent;padding:7px 12px;border-radius:6px;cursor:pointer;font-size:13px;display:inline-flex;align-items:center}
.tabs button.active{background:#202633;color:#fff;border-color:#2f3d55}
main{padding:20px}
table{width:100%;border-collapse:collapse;font-size:13px}
th,td{text-align:left;padding:8px 10px;border-bottom:1px solid #20242e}
th{color:#8b93a7;font-weight:500;font-size:12px;text-transform:uppercase;user-select:none}
th.sortable{cursor:pointer}
th.sortable:hover{color:#fff;background:#1c222b}
tr:hover td{background:#161a22}
.pill{display:inline-block;padding:2px 8px;border-radius:10px;font-size:11px;background:#232a37;color:#aab}
.pill.ok{background:#1e3a26;color:#71d68a}
.pill.warn{background:#3a3520;color:#e0c66a}
.pill.err{background:#3a1f1f;color:#e07a7a}
.pill.chat{background:#2a1f3a;color:#c58af9}
.pill.rubric{background:#3a2f1f;color:#e0c66a}
.btn{background:#2a3340;color:#e6e6e6;border:1px solid #334055;padding:5px 10px;border-radius:5px;cursor:pointer;font-size:12px}
.btn:hover{background:#334055}
.btn.danger{background:#3a1f1f;border-color:#7f1d1d}
.btn.green{background:#1e3a26;border-color:#2d5a3a;color:#71d68a}
.btn.blue{background:#1e40af;border-color:#1e88e5;color:#fff}
.row-flex{display:flex;gap:8px;flex-wrap:wrap;align-items:center;margin-bottom:12px}
input,select,textarea{background:#171a21;color:#e6e6e6;border:1px solid #2a2f3a;padding:6px 10px;border-radius:5px;font-size:13px;max-width:280px;font-family:inherit}
.modal{position:fixed;inset:0;background:rgba(0,0,0,.8);display:none;align-items:center;justify-content:center;z-index:50}
.modal img{max-width:92vw;max-height:92vh;border-radius:8px}
.modal.on{display:flex}
.hint{color:#8b93a7;font-size:12px}
.kpi{display:flex;gap:12px;flex-wrap:wrap;margin-bottom:16px}
.kpi div{background:#151821;border:1px solid #232a37;border-radius:8px;padding:10px 16px;min-width:140px}
.kpi b{display:block;font-size:20px;color:#8ab4f8;margin-top:2px}
.recv-wrap{max-width:420px;font-size:12px}
.recv-visible{color:#aab}
.recv-hidden{display:none;color:#aab}
.recv-toggle{color:#7cc4ff;cursor:pointer;font-size:11px;margin-left:4px;text-decoration:underline}
tr.rubric-row td{background:#2a2410 !important}
.vault-card{background:#151821;border:1px solid #232a37;border-radius:8px;padding:14px 18px;margin-bottom:12px}
.vault-card h3{margin:0 0 10px;font-size:13px;color:#8ab4f8;text-transform:uppercase}
.calendar-grid{display:grid;grid-template-columns:repeat(7,1fr);gap:6px;max-width:560px}
.cal-day{background:#10131a;border:1px solid #232a37;border-radius:6px;padding:6px;text-align:center;cursor:pointer;min-height:56px;font-size:11px}
.cal-day:hover{border-color:#4a7ce0}
.cal-day .num{font-size:14px;font-weight:600;color:#e6e6e6}
.cal-day .cnt{color:#8ab4f8;font-size:11px;margin-top:2px}
.cal-day .rain{color:#71d68a;font-size:11px}
/* === Промокоды === */
.promo-card{background:#151821;border:1px solid #232a37;border-radius:8px;padding:12px 16px;margin-bottom:10px}
.promo-card.used{opacity:.55;border-color:#2a2f3a}
.promo-head{display:flex;justify-content:space-between;align-items:center;margin-bottom:8px;gap:8px;flex-wrap:wrap}
.promo-date{color:#8b93a7;font-size:12px}
.promo-text{white-space:pre-wrap;font-size:13px;line-height:1.5;color:#e6e6e6;margin-bottom:10px;font-family:inherit}
.promo-actions{display:flex;gap:8px;flex-wrap:wrap;align-items:center}
.promo-actions a{color:#7cc4ff;font-size:12px;text-decoration:none}
.promo-actions a:hover{text-decoration:underline}

/* === Языки / Игроки === */
.lang-grid{display:grid;grid-template-columns:repeat(auto-fill,minmax(240px,1fr));gap:12px}
.lang-card{background:#151821;border:1px solid #232a37;border-radius:8px;padding:12px 16px}
.lang-card h4{margin:0 0 6px;color:#fff;font-size:14px}
.player-row{padding:3px 0;color:#aab;font-size:12px;border-top:1px solid #1e2430}
.player-row:first-of-type{border-top:none;margin-top:6px;padding-top:6px}

/* === AI === */
.ai-grid{display:grid;grid-template-columns:1fr 1fr;gap:16px}
.ai-col h3{margin:0 0 8px;font-size:13px;color:#8ab4f8;text-transform:uppercase}
.ai-msg,.ai-phrase{padding:8px 12px;background:#151821;border:1px solid #232a37;border-radius:6px;margin-bottom:6px;cursor:pointer;font-size:13px;user-select:none}
.ai-msg:hover,.ai-phrase:hover{border-color:#4a7ce0;background:#1a1f2b}
.ai-msg.copied,.ai-phrase.copied{border-color:#71d68a;background:#162a1e}

/* === Ставки (чатовые) === */
tr.bet-row.chat td{background:rgba(197,138,249,.06)}
tr.bet-row.chat:hover td{background:rgba(197,138,249,.12)}

/* === Модалка скриншота === */
.modal{position:fixed;inset:0;background:rgba(0,0,0,.85);display:none;align-items:center;justify-content:center;z-index:99;cursor:zoom-out}
.modal.on{display:flex}
.modal img{max-width:92vw;max-height:92vh;border-radius:8px;box-shadow:0 8px 40px rgba(0,0,0,.6)}
</style></head><body>
<header>
  <h1>🎰 Shuffle Monitor</h1>
  <span id="uptime" class="hint">—</span>
  <a href="/" style="margin-left:16px">← На сайт</a>
  <span style="margin-left:auto" class="pill" id="kpiTotal">—</span>
  <div class="row-flex" style="margin:0">
    <button class="btn green" onclick="controlAction('start')" title="Запуск">▶️</button>
    <button class="btn danger" onclick="controlAction('restart')" title="Перезапуск">🔄</button>
    <button class="btn danger" onclick="controlAction('stop')" title="Стоп">⛔</button>
  </div>
  <span class="pill" id="ngrokUrl" style="cursor:pointer" onclick="copyNgrok()">🌐 ngrok: —</span>
  <span class="pill" id="deployInfo" title="Деплой">🚀 —</span>
</header>
<div class="tabs">
  <button data-tab="status" class="active">Вкладки</button>
  <button data-tab="events">События</button>
  <button data-tab="bulk">Bulk</button>
  <button data-tab="bets">Ставки</button>
  <button data-tab="betsrains">Ставки→Дожди</button>
  <button data-tab="contexts">Контексты</button>
  <button data-tab="intents">Интенты</button>
  <button data-tab="promos">Промокоды</button>
  <button data-tab="calendar">📅 Календарь</button>
  <button data-tab="chart">📈 Графики</button>
  <button data-tab="activity">📊 Активность</button>
  <button data-tab="mods">🛡 Модераторы</button>
  <button data-tab="me">👤 Профиль</button>
  <button data-tab="ai">🤖 AI</button>
  <button data-tab="players">👥 Игроки</button>
  <button data-tab="langs">🌍 Языки</button>
  <button data-tab="db">🗄 БД</button>
</div>
<main>
  <div id="tab-status"><div class="kpi" id="kpiRow"></div>
    <table id="tbl-status"><thead><tr><th>#</th><th>Режим</th><th>Чат</th><th>Статус</th><th>Обновлено</th><th>📸</th></tr></thead>
    <tbody id="instBody"></tbody></table></div>

  <div id="tab-events" style="display:none">
    <div class="row-flex">
      <input type="text" id="fSender" placeholder="Отправитель">
      <select id="fChat"><option value="">Все чаты</option></select>
      <select id="fType"><option value="">Все</option><option value="tip">Tip</option><option value="rain">Rain</option></select>
      <select id="fRange"><option value="">Всё</option><option value="1">1 час</option><option value="24" selected>24 часа</option><option value="168">7 дней</option><option value="720">Месяц</option></select>
      <button class="btn" onclick="loadEvents()">Фильтр</button>
    </div>
    <table id="tbl-events"><thead><tr>
      <th class="sortable" onclick="sortTable('tbl-events',0)">Время</th>
      <th class="sortable" onclick="sortTable('tbl-events',1)">Тип</th>
      <th class="sortable" onclick="sortTable('tbl-events',2)">Чат</th>
      <th class="sortable" onclick="sortTable('tbl-events',3)">От</th>
      <th>Кому</th>
      <th class="sortable" onclick="sortTable('tbl-events',5)">Сумма</th>
      <th class="sortable" onclick="sortTable('tbl-events',6)">₽</th>
      <th class="sortable" onclick="sortTable('tbl-events',7)">₽/чел</th>
      <th class="sortable" onclick="sortTable('tbl-events',8)">$</th>
    </tr></thead>
    <tbody id="evBody"></tbody></table></div>

  <div id="tab-bets" style="display:none">
    <div class="row-flex">
      <input type="text" id="betSender" placeholder="Отправитель">
      <select id="betChat"><option value="">Все чаты</option></select>
      <select id="betSource"><option value="">Все</option><option value="bets_page">Таблица</option><option value="chat">Чат</option></select>
      <input type="number" id="betMinUsd" placeholder="Мин. $" value="50">
      <select id="betRange"><option value="">Всё</option><option value="1">1 час</option><option value="24" selected>24 часа</option><option value="168">7 дней</option><option value="720">Месяц</option></select>
      <button class="btn" onclick="loadBets()">Фильтр</button>
      <button class="btn" onclick="loadBetsTop()">🏆 Топ</button>
      <button class="btn danger" onclick="clearBetsTop()">🗑 Очистить топ</button>
      <button class="btn danger" onclick="clearBetsTopHard()">💥 Удалить всё</button>
    </div>
    <div id="betStats" style="margin-bottom:16px"></div>
    <table id="tbl-bets"><thead><tr>
      <th class="sortable" onclick="sortTable('tbl-bets',0)">Время</th>
      <th>Источник</th>
      <th class="sortable" onclick="sortTable('tbl-bets',2)">Чат</th>
      <th class="sortable" onclick="sortTable('tbl-bets',3)">Игрок</th>
      <th>Игра</th>
      <th>Множ.</th>
      <th>Сумма</th>
      <th class="sortable" onclick="sortTable('tbl-bets',7)">$</th>
      <th class="sortable" onclick="sortTable('tbl-bets',8)">₽</th>
    </tr></thead>
    <tbody id="betsBody"></tbody></table></div>

  <div id="tab-betsrains" style="display:none">
    <div class="row-flex">
      <input type="number" id="b2rWindow" value="900" style="width:100px">
      <span class="hint">сек — окно до дождя</span>
      <button class="btn" onclick="loadBetsToRains()">Обновить</button>
      <button class="btn danger" onclick="clearBetsToRains()">🗑 Очистить</button>
    </div>
    <div id="b2rStats" style="margin-bottom:16px"></div>
    <table id="tbl-b2r"><thead><tr>
      <th>#</th>
      <th class="sortable" onclick="sortTable('tbl-b2r',1)">Ник</th>
      <th class="sortable" onclick="sortTable('tbl-b2r',2)">Ставок ≥$20</th>
      <th class="sortable" onclick="sortTable('tbl-b2r',3)">Привели к дождю</th>
      <th class="sortable" onclick="sortTable('tbl-b2r',4)">%</th>
    </tr></thead>
    <tbody id="b2rBody"></tbody></table></div>

  <div id="tab-bulk" style="display:none">
    <table id="tbl-bulk"><thead><tr>
      <th class="sortable" onclick="sortTable('tbl-bulk',0)">Время</th>
      <th class="sortable" onclick="sortTable('tbl-bulk',1)">Чат</th>
      <th class="sortable" onclick="sortTable('tbl-bulk',2)">От</th>
      <th>Сумма</th>
      <th class="sortable" onclick="sortTable('tbl-bulk',4)">₽/чел</th>
      <th class="sortable" onclick="sortTable('tbl-bulk',5)">₽ всего</th>
      <th class="sortable" onclick="sortTable('tbl-bulk',6)">Получателей</th>
      <th>Кому</th>
      <th>Действие</th>
    </tr></thead>
    <tbody id="bulkBody"></tbody></table></div>

  <div id="tab-contexts" style="display:none">
    <div class="row-flex">
      <select id="ctxChat"><option value="">Все чаты</option></select>
      <input type="number" id="ctxLimit" value="30" style="width:80px">
      <button class="btn" onclick="loadContexts()">Обновить</button>
    </div>
    <div id="ctxList"></div></div>

  <div id="tab-intents" style="display:none">
    <div class="row-flex">
      <select id="intChat"><option value="">Все чаты</option></select>
      <input type="text" id="intSender" placeholder="Игрок">
      <select id="intOnly"><option value="">Все</option><option value="1">Сбывшиеся</option><option value="0">Несбывшиеся</option></select>
      <button class="btn" onclick="loadIntents()">Фильтр</button>
      <button class="btn" onclick="loadIntentsStats()">🏆 Рейтинг</button>
    </div>
    <div id="intStats" style="margin-bottom:16px"></div>
    <table id="tbl-intents"><thead><tr>
      <th class="sortable" onclick="sortTable('tbl-intents',0)">Время</th>
      <th class="sortable" onclick="sortTable('tbl-intents',1)">Чат</th>
      <th class="sortable" onclick="sortTable('tbl-intents',2)">Игрок</th>
      <th>Текст</th>
      <th class="sortable" onclick="sortTable('tbl-intents',4)">AI%</th>
      <th>Статус</th>
      <th class="sortable" onclick="sortTable('tbl-intents',6)">Δ сек</th>
    </tr></thead>
    <tbody id="intBody"></tbody></table></div>

<div id="tab-promos" style="display:none">
  <div class="row-flex">
    <select id="promoFilter" onchange="loadPromos()">
      <option value="">Все</option>
      <option value="0">Неиспользованные</option>
      <option value="1">Использованные</option>
    </select>
    <button class="btn" onclick="loadPromos()">Обновить</button>
    <button class="btn" onclick="testPromo(event)">🔍 Проверить канал</button>
    <span class="hint" id="promoStats"></span>
  </div>
  <div id="promoList"></div>
  
</div>
  <div id="tab-calendar" style="display:none">
    <div class="vault-card">
      <h3>📅 Календарь</h3>
      <div class="row-flex">
        <input type="date" id="docxDate">
        <button class="btn green" onclick="downloadDocx()">📄 Скачать docx</button>
      </div>
      <div class="calendar-grid" id="calGrid"></div>
    </div>
    <div id="calDetail"></div>
  </div>

  <div id="tab-chart" style="display:none">
    <div class="row-flex">
      <select id="cType"><option value="">Всё</option><option value="tip">Tip</option><option value="rain">Rain</option></select>
      <select id="cChat"><option value="">Все чаты</option></select>
      <select id="cRange"><option value="">Всё время</option><option value="1">1 час</option><option value="24" selected>24 часа</option><option value="168">7 дней</option><option value="720">Месяц</option></select>
      <button class="btn" onclick="loadChart()">Обновить</button>
    </div>
    <div style="background:#151821;border:1px solid #232a37;border-radius:8px;padding:16px;margin-bottom:16px"><canvas id="hourCount" height="90"></canvas></div>
    <div style="background:#151821;border:1px solid #232a37;border-radius:8px;padding:16px"><canvas id="hourSum" height="90"></canvas></div>
  </div>

  <div id="tab-activity" style="display:none">
    <div class="row-flex">
      <label class="hint">Окно: </label>
      <select id="actWindow" onchange="loadActivityDashboard()">
        <option value="60" selected>60 мин</option>
        <option value="120">2 часа</option>
        <option value="1440">24 часа</option>
      </select>
      <button class="btn" onclick="loadActivityDashboard()">Обновить</button>
      <button class="btn blue" onclick="loadRainActivity()">🌧 Rain Activity</button>
      <span class="hint" id="actStatus"></span>
    </div>
    <div id="actGrid"></div>
    <div id="rainActivityBox"></div>
  </div>

  <div id="tab-mods" style="display:none">
    <div class="row-flex">
      <input type="text" id="modSearch" placeholder="Поиск по нику">
      <button class="btn" onclick="loadModerators()">Обновить</button>
      <span class="hint" id="modStats"></span>
    </div>
    <table id="tbl-mods"><thead><tr>
      <th>Ник</th><th>Первый раз</th><th>Последний раз</th><th>Чаты</th><th>Действие</th>
    </tr></thead>
    <tbody id="modBody"></tbody></table>
    <div id="modMessages"></div>
  </div>

  <div id="tab-me" style="display:none">
    <div class="vault-card">
      <div class="row-flex" style="justify-content:space-between;margin-bottom:0">
        <h3 style="margin:0">👤 Мой профиль</h3>
        <div class="row-flex" style="margin:0">
          <input type="text" id="meNick" placeholder="Ник" style="width:200px">
          <button class="btn green" onclick="loadMe()">🔎 Проверить</button>
          <button class="btn" onclick="saveMeNick()">💾 Запомнить</button>
          <span class="hint" id="meStatus"></span>
        </div>
      </div>
      <div id="meTop"></div>
    </div>
    <div class="vault-card"><h3>📊 Активность</h3><div id="meActivity"></div></div>
    <div class="vault-card"><h3>🏆 Мои дожди</h3>
      <div class="row-flex">
        <select id="meRainDays" onchange="loadMeRains()" style="max-width:160px">
          <option value="7">7 дней</option><option value="30" selected>30 дней</option><option value="90">90 дней</option>
        </select>
        <span class="hint" id="meRainsStatus"></span>
      </div>
      <div id="meRains"></div>
    </div>
    <div class="vault-card"><h3>🎰 Мои ставки</h3><div id="meBets"></div></div>
  </div>

  <div id="tab-ai" style="display:none">
    <div class="row-flex">
      <select id="aiChat"><option value="">Выбери чат</option></select>
      <button class="btn" onclick="runAI()">🤖 Получить</button>
      <button class="btn" onclick="loadAILatest()">📋 Последний</button>
      <span class="hint" id="aiStatus"></span>
    </div>
    <div id="aiTopic"></div>
    <div id="aiResult"></div></div>

  <div id="tab-players" style="display:none">
    <div class="row-flex">
      <select id="plChatFilter"><option value="">Все чаты</option></select>
      <input type="text" id="plUserFilter" placeholder="Ник">
      <label style="font-size:12px;color:#aab"><input type="checkbox" id="plOnlineOnly"> Онлайн</label>
      <button class="btn" onclick="loadPlayersTab()">Обновить</button>
    </div>
    <div id="plChatsSummary" style="margin-bottom:16px"></div>
    <table id="tbl-players"><thead><tr>
      <th class="sortable" onclick="sortTable('tbl-players',0)">Чат</th>
      <th class="sortable" onclick="sortTable('tbl-players',1)">Игрок</th>
      <th class="sortable" onclick="sortTable('tbl-players',2)">Сообщений</th>
      <th class="sortable" onclick="sortTable('tbl-players',3)">Последний раз</th>
      <th>Статус</th><th>Действие</th>
    </tr></thead>
    <tbody id="plPlayersBody"></tbody></table></div>

  <div id="tab-langs" style="display:none">
    <div class="row-flex">
      <input type="text" id="langFilter" placeholder="Поиск по нику">
      <button class="btn" onclick="loadLanguages()">Обновить</button>
    </div>
    <div id="langGrid" class="lang-grid"></div></div>

  <div id="tab-db" style="display:none">
    <div class="vault-card">
      <h3>🚀 Деплой</h3>
      <div id="dbDeploy">—</div>
    </div>
    <div class="vault-card">
      <h3>🗄 База данных</h3>
      <div id="dbStats">—</div>
    </div>
  </div>
</main>

<script>
const $ = id => document.getElementById(id);
const fmtTime = ts => new Date(ts*1000).toLocaleString('ru-RU');
const fmtMoney = v => (v||0).toLocaleString('ru-RU',{maximumFractionDigits:2});
const esc = s => (s==null?'':String(s)).replace(/[<>&"]/g, c=>({'<':'&lt;','>':'&gt;','&':'&amp;','"':'&quot;'}[c]));

function sortTable(tableId, colIdx) {
  const table = document.getElementById(tableId);
  if (!table) return;
  const tb = table.querySelector('tbody');
  const th = table.querySelectorAll('th')[colIdx];
  const dir = (th.dataset.dir === 'asc') ? -1 : 1;
  th.dataset.dir = dir === 1 ? 'asc' : 'desc';
  const rows = Array.from(tb.querySelectorAll('tr'));
  rows.sort((a,b) => {
    const av = (a.cells[colIdx]?.getAttribute('data-v') ?? a.cells[colIdx]?.textContent ?? '').trim();
    const bv = (b.cells[colIdx]?.getAttribute('data-v') ?? b.cells[colIdx]?.textContent ?? '').trim();
    const an = parseFloat(av), bn = parseFloat(bv);
    if (!isNaN(an) && !isNaN(bn)) return dir * (an - bn);
    return dir * av.localeCompare(bv);
  });
  rows.forEach(r => tb.appendChild(r));
}

function renderReceivers(recs, collapseAfter) {
  recs = recs || [];
  if (!recs.length) return '<span class="hint">—</span>';
  const n = collapseAfter || 5;
  if (recs.length <= n) return '<span class="recv-visible">'+recs.map(esc).join(', ')+'</span>';
  const id = 'recv_' + Math.random().toString(36).slice(2,9);
  return '<span class="recv-visible">'+recs.slice(0,n).map(esc).join(', ')+'</span>'
    + '<span class="recv-hidden" id="'+id+'">, '+recs.slice(n).map(esc).join(', ')+'</span>'
    + ' <span class="recv-toggle" onclick="document.getElementById(\''+id+'\').style.display=\'inline\';this.style.display=\'none\'">+'+(recs.length-n)+' ещё</span>';
}

function switchTab(name){
  document.querySelectorAll('.tabs button').forEach(b=>b.classList.toggle('active', b.dataset.tab===name));
  ['status','events','bulk','bets','betsrains','contexts','intents','promos','calendar','chart','ai','players','langs','me','mods','activity','db'].forEach(t=>{
    const el = $('tab-'+t); if(el) el.style.display = (t===name) ? '' : 'none';
  });
  if(name==='events') { loadChats(); loadEvents(); }
  if(name==='bulk') loadBulk();
  if(name==='bets') { loadChats(); loadBets(); }
  if(name==='betsrains') loadBetsToRains();
  if(name==='contexts') { loadChats(); loadContexts(); }
  if(name==='intents') { loadChats(); loadIntents(); }
  if(name==='promos') loadPromos();
  if(name==='calendar') loadCalendar();
  if(name==='chart') { loadChats(); loadChart(); }
  if(name==='me') loadMe();
  if(name==='ai') { loadChats(); loadAILatest(); }
  if(name==='players') { loadChats(); loadPlayersTab(); }
  if(name==='langs') loadLanguages();
  if(name==='mods') loadModerators();
  if(name==='activity') loadActivityDashboard();
  if(name==='db') loadDb();
}
document.querySelectorAll('.tabs button').forEach(b=>b.onclick=()=>switchTab(b.dataset.tab));

async function refreshStatus(){
  try {
    const r = await fetch('/api/status').then(r=>r.json());
    $('uptime').textContent = r.first_ts ? 'с ' + new Date(r.first_ts*1000).toLocaleDateString('ru-RU') : '—';
    $('kpiTotal').textContent = `Всего: ${r.total||0} | Tips: ${r.tips||0} | Rains: ${r.rains||0} | Bets: ${r.bets||0} | Intents: ${r.intents||0}`;
    const now = Date.now()/1000;
    $('kpiRow').innerHTML = '<div>Всего<b>'+(r.total||0)+'</b></div><div>Tips<b>'+(r.tips||0)+'</b></div><div>Rains<b>'+(r.rains||0)+'</b></div><div>Bets<b>'+(r.bets||0)+'</b></div><div>Intents<b>'+(r.intents||0)+'</b></div><div>Онлайн<b>'+(r.instances||[]).filter(i=>now-(i.updated_at||0)<90).length+'</b></div>';
    const body = $('instBody'); body.innerHTML = '';
    (r.instances||[]).forEach(i=>{
      const online = now - (i.updated_at||0) < 90;
      let cls='pill', txt=i.status||'—';
      if(!online){ cls+=' err'; txt='offline'; }
      else if(i.status==='ok') cls+=' ok';
      else if(i.status==='waiting'||i.status==='rotating'||i.status==='busy') cls+=' warn';
      else if(['error','timeout','verify_error','chat_not_loaded'].includes(i.status)) cls+=' err';
      const upd = i.updated_at ? Math.round(now - i.updated_at)+'с' : '—';
      let chatsCell = '<b>' + esc(i.current_chat||'—') + '</b>';
      const chats = i.chats || [];
      if (i.mode === 'rotation' && chats.length > 0) {
        chatsCell = chats.map(c => c === i.current_chat
          ? '<b style="color:#71d68a">' + esc(c) + '</b>'
          : '<span class="hint">' + esc(c) + '</span>'
        ).join(' → ');
      }
      const tr = document.createElement('tr');
      tr.innerHTML = '<td>'+i.iid+'</td><td>'+esc(i.mode||'')+'</td><td>'+chatsCell+'</td><td><span class="'+cls+'">'+esc(txt)+'</span></td><td>'+upd+'</td><td><button class="btn" onclick="reqShot('+i.iid+')">📸</button></td>';
      body.appendChild(tr);
    });
  } catch(e){}
}
async function reqShot(iid){
  await fetch('/api/screenshot/'+iid,{method:'POST'});
  let t=0; while(t++<40){ await new Promise(r=>setTimeout(r,1500));
    const r = await fetch('/api/screenshot/'+iid);
    if(r.status===200){ $('imgModalImg').src = URL.createObjectURL(await r.blob()); $('imgModal').classList.add('on'); return; }
    if(r.status===404){ alert('Не удалось'); return; }
  } alert('Таймаут');
}
async function loadChats(){
  const r = await fetch('/api/chats').then(r=>r.json());
  ['fChat','betChat','ctxChat','intChat','aiChat','plChatFilter','cChat'].forEach(id=>{
    const sel = $(id); if(!sel) return; const cur = sel.value;
    const ph = id==='aiChat' ? '<option value="">Выбери чат</option>' : '<option value="">Все чаты</option>';
    sel.innerHTML = ph + r.chats.map(c=>'<option value="'+esc(c)+'">'+esc(c)+'</option>').join('');
    sel.value = cur;
  });
}
async function loadEvents(){
  const p = new URLSearchParams();
  if($('fSender').value) p.set('sender',$('fSender').value);
  if($('fChat').value) p.set('chat',$('fChat').value);
  if($('fType').value) p.set('type',$('fType').value);
  if($('fRange').value) p.set('last_hours',$('fRange').value);
  const r = await fetch('/api/events?'+p).then(r=>r.json());
  const body = $('evBody'); body.innerHTML = '';
  r.events.forEach((e,i)=>{
    const tr = document.createElement('tr'); const per = e.per_person_rub || 0;
    tr.innerHTML = '<td data-v="'+(e.ts||0)+'">'+(e.ts?fmtTime(e.ts):esc(e.created_at))+'</td>'
      + '<td data-v="'+esc(e.type)+'"><span class="pill">'+esc(e.type)+'</span></td>'
      + '<td data-v="'+esc(e.chat)+'">'+esc(e.chat)+'</td>'
      + '<td data-v="'+esc(e.sender)+'">'+esc(e.sender)+'</td>'
      + '<td class="recv-wrap">'+renderReceivers(e.receivers)+'</td>'
      + '<td data-v="'+esc(e.amount_text)+'">'+esc(e.amount_text)+'</td>'
      + '<td data-v="'+(e.amount_rub||0)+'">'+fmtMoney(e.amount_rub)+'</td>'
      + '<td data-v="'+per+'" style="color:#8ab4f8">'+(per?fmtMoney(per):'—')+'</td>'
      + '<td data-v="'+(e.amount_usd||0)+'" style="color:#71d68a">$'+fmtMoney(e.amount_usd||0)+'</td>';
    body.appendChild(tr);
  });
}
async function loadBulk(){
  const r = await fetch('/api/bulk').then(r=>r.json());
  const body = $('bulkBody'); body.innerHTML='';
  if(!r.alerts.length){ body.innerHTML = '<tr><td colspan="9" class="hint">Нет данных</td></tr>'; return; }
  r.alerts.forEach(a=>{
    const recCount = (a.receivers||[]).length; const rubEach = a.amount_rub || 0;
    const isRubric = !!a.is_rubric;
    const tr = document.createElement('tr'); if(isRubric) tr.className = 'rubric-row';
    tr.innerHTML = '<td data-v="'+(a.ts||0)+'">'+fmtTime(a.ts)+'</td>'
      + '<td data-v="'+esc(a.chat)+'">'+esc(a.chat)+(isRubric?' <span class="pill rubric">📌 РУБРИКА</span>':'')+'</td>'
      + '<td data-v="'+esc(a.sender)+'">'+esc(a.sender)+'</td>'
      + '<td data-v="'+esc(a.amount_text)+'">'+esc(a.amount_text)+'</td>'
      + '<td data-v="'+rubEach+'">'+fmtMoney(rubEach)+' ₽</td>'
      + '<td data-v="'+(rubEach*recCount)+'" style="color:#71d68a">'+fmtMoney(rubEach*recCount)+' ₽</td>'
      + '<td data-v="'+recCount+'">'+recCount+'</td>'
      + '<td class="recv-wrap">'+renderReceivers(a.receivers)+'</td>'
      + '<td><button class="btn" onclick="toggleRubric('+a.id+')">'+(isRubric?'✓ Снять':'📌 Рубрика')+'</button></td>';
    body.appendChild(tr);
  });
}
async function toggleRubric(id){ await fetch('/api/bulk/rubric/'+id,{method:'POST'}); loadBulk(); }

async function loadBets(){
  const p = new URLSearchParams();
  if($('betSender').value) p.set('sender',$('betSender').value);
  if($('betChat').value) p.set('chat',$('betChat').value);
  if($('betSource').value) p.set('source',$('betSource').value);
  if($('betMinUsd').value) p.set('min_usd',$('betMinUsd').value);
  if($('betRange').value) p.set('last_hours',$('betRange').value);
  const r = await fetch('/api/bets?'+p).then(r=>r.json());
  const body = $('betsBody'); body.innerHTML='';
  if(!r.bets.length){ body.innerHTML = '<tr><td colspan="9" class="hint">Нет ставок · query='+esc(p.toString())+'</td></tr>'; return; }
  r.bets.forEach(b=>{
    const fromChat = b.source === 'chat';
    const tr = document.createElement('tr'); tr.className = 'bet-row' + (fromChat?' chat':'');
    const realInfo = b.real_sender ? ' <span class="pill ok">real: '+esc(b.real_sender)+'</span>' : '';
    tr.innerHTML = '<td data-v="'+(b.ts||0)+'">'+fmtTime(b.ts)+'</td>'
      + '<td><span class="pill '+(fromChat?'chat':'')+'">'+(fromChat?'💬':'📊')+'</span></td>'
      + '<td data-v="'+esc(b.chat)+'">'+esc(b.chat)+'</td>'
      + '<td data-v="'+esc(b.sender)+'"><b>'+esc(b.sender)+'</b>'+realInfo+'</td>'
      + '<td>'+esc(b.game)+'</td><td>'+esc(b.multiplier)+'</td>'
      + '<td>'+esc(b.amount_text)+' '+esc(b.crypto)+'</td>'
      + '<td data-v="'+(b.amount_usd||0)+'" style="color:#71d68a">$'+fmtMoney(b.amount_usd)+'</td>'
      + '<td data-v="'+(b.amount_rub||0)+'">'+fmtMoney(b.amount_rub)+'</td>';
    body.appendChild(tr);
  });
}
async function loadBetsTop(){
  const p = new URLSearchParams();
  if($('betRange').value) p.set('last_hours',$('betRange').value);
  const r = await fetch('/api/bets_top_senders?'+p).then(r=>r.json());
  const el = $('betStats');
  if(!r.senders.length){ el.innerHTML = '<p class="hint">Пусто (после очистки топ обнулён)</p>'; return; }
  el.innerHTML = '<div class="kpi">' + r.senders.map(s=>'<div>'+esc(s.nickname)+'<b>'+s.cnt+' шт.</b><span class="hint">$'+fmtMoney(s.total_usd)+'</span></div>').join('') + '</div>';
}
async function clearBetsTop(){
  if(!confirm('Очистить топ отправителей? (записи ставок останутся, но топ будет считать только новые)')) return;
  const r = await fetch('/api/bets/clear_top',{method:'POST'}).then(r=>r.json());
  if(r.ok){ alert('✅ Топ очищен (timestamp: '+new Date(r.cleared_at*1000).toLocaleTimeString('ru-RU')+')'); loadBetsTop(); }
  else alert('❌ Ошибка');
}
async function clearBetsTopHard(){
  if(!confirm('⚠️ УДАЛИТЬ ВСЕ ставки из БД? Это необратимо!')) return;
  if(!confirm('Точно? Все данные таблицы bets будут удалены.')) return;
  const r = await fetch('/api/bets/clear_top?hard=1',{method:'POST'}).then(r=>r.json());
  if(r.ok){ alert('✅ Удалено записей: '+r.deleted); loadBets(); loadBetsTop(); }
  else alert('❌ Ошибка');
}
async function clearBetsToRains(){
  if(!confirm('Очистить историю Bets→Rains?')) return;
  await fetch('/api/bets_to_rains/clear',{method:'POST'}).then(r=>r.json());
  alert('✅ Очищено');
  loadBetsToRains();
}
async function loadBetsToRains(){
  const w = $('b2rWindow').value || 900;
  const r = await fetch('/api/bets_to_rains?window_sec='+w).then(r=>r.json());
  const body = $('b2rBody'); body.innerHTML='';
  $('b2rStats').innerHTML = '<div class="kpi"><div>Окно<b>'+r.window_sec+'с</b></div><div>Игроков<b>'+r.stats.length+'</b></div></div>';
  if(!r.stats.length){ body.innerHTML = '<tr><td colspan="5" class="hint">Нет данных (нужно 3+ ставки ≥$20)</td></tr>'; return; }
  r.stats.forEach((s,i)=>{ const cls = s.percent>=50?'ok':(s.percent>=25?'warn':'err');
    const tr = document.createElement('tr');
    tr.innerHTML = '<td>'+(i+1)+'</td><td data-v="'+esc(s.nickname)+'"><b>'+esc(s.nickname)+'</b></td><td data-v="'+s.bets+'">'+s.bets+'</td><td data-v="'+s.triggered+'">'+s.triggered+'</td><td data-v="'+s.percent+'"><span class="pill '+cls+'">'+s.percent+'%</span></td>';
    body.appendChild(tr);
  });
}
async function loadContexts(){
  const p = new URLSearchParams();
  if($('ctxChat').value) p.set('chat',$('ctxChat').value);
  if($('ctxLimit').value) p.set('limit',$('ctxLimit').value);
  const r = await fetch('/api/contexts?'+p).then(r=>r.json());
  const el = $('ctxList');
  if(!r.contexts.length){ el.innerHTML = '<p class="hint">Нет</p>'; return; }
  el.innerHTML = r.contexts.map(c=>{
    const rub = c.amount_rub ? fmtMoney(c.amount_rub) + ' ₽' : '';
    const msgs = (c.context_messages||[]).slice(-15); const bets = (c.context_bets||[]).slice(-5);
    const mH = msgs.length ? msgs.map((m,i)=>'<div style="padding:4px 8px;border-left:3px solid #2a2f3a;margin:4px 0;color:#aab;font-size:12px">'+(i+1)+'. '+esc(m)+'</div>').join('') : '<div class="hint">—</div>';
    const bH = bets.length ? bets.map(b=>{
      const amt = esc(b.amount_text||'');
      const sym = esc(b.crypto_symbol||'');
      const rub2 = (b.amount_rub||0) > 0 ? ' <span style="color:#4ade80">~'+Math.round(b.amount_rub)+'₽</span>' : '';
      const usd2 = (b.amount_usd||0) > 0 ? ' <span style="color:#8ab4f8">$'+b.amount_usd.toFixed(2)+'</span>' : '';
      return '<div style="padding:4px 8px;border-left:3px solid #71d68a;margin:4px 0;color:#71d68a;font-size:12px">🎰 <b>'+esc(b.sender||'?')+'</b> · '+esc(b.game||'')+' · '+esc(b.multiplier||'')+' · '+amt+' '+sym+rub2+usd2+'</div>';
    }).join('') : '<div class="hint">—</div>';
    return '<div style="background:#151821;border:1px solid #232a37;border-radius:8px;padding:14px;margin-bottom:12px">'
      + '<div style="display:flex;gap:12px;flex-wrap:wrap;align-items:center;margin-bottom:10px">'
      + '<span class="pill">🌧 RAIN</span><b>'+esc(c.chat)+'</b>'
      + '<span class="hint">'+(c.ts?fmtTime(c.ts):esc(c.created_at))+'</span>'
      + '<span style="margin-left:auto;color:#71d68a;font-weight:bold">'+esc(c.amount_text)+' '+esc(c.crypto_symbol||'')+' '+(rub?'· '+rub:'')+'</span></div>'
      + '<div style="color:#aab;font-size:12px;margin-bottom:6px">👤 <b>'+esc(c.sender)+'</b> → '+(c.receivers||[]).length+'</div>'
      + '<div class="recv-wrap" style="margin-bottom:8px">'+renderReceivers(c.receivers,8)+'</div>'
      + '<div style="display:grid;grid-template-columns:1fr 1fr;gap:12px"><div>'+mH+'</div><div>'+bH+'</div></div></div>';
  }).join('');
}
async function loadIntents(){
  const p = new URLSearchParams();
  if($('intChat').value) p.set('chat',$('intChat').value);
  if($('intSender').value) p.set('sender',$('intSender').value);
  if($('intOnly').value !== '') p.set('only_matched',$('intOnly').value);
  const r = await fetch('/api/intents?'+p).then(r=>r.json());
  const body = $('intBody'); body.innerHTML='';
  if(!r.intents.length){ body.innerHTML = '<tr><td colspan="7" class="hint">Нет данных</td></tr>'; return; }
  r.intents.forEach(x=>{
    const matched = !!x.matched_rain_id;
    const tr = document.createElement('tr');
    tr.innerHTML = '<td data-v="'+(x.ts||0)+'">'+fmtTime(x.ts)+'</td>'
      + '<td data-v="'+esc(x.chat)+'">'+esc(x.chat)+'</td>'
      + '<td data-v="'+esc(x.sender)+'"><b>'+esc(x.sender)+'</b></td>'
      + '<td>'+esc(x.text)+'</td>'
      + '<td data-v="'+(x.percent||0)+'"><span class="pill '+(x.percent>=70?'ok':'warn')+'">'+(x.percent||100)+'%</span></td>'
      + '<td><span class="pill '+(matched?'ok':'err')+'">'+(matched?'✅':'❌')+'</span></td>'
      + '<td data-v="'+(x.matched_delta_sec||0)+'">'+(x.matched_delta_sec ? Math.round(x.matched_delta_sec)+'с' : '—')+'</td>';
    body.appendChild(tr);
  });
}
async function loadIntentsStats(){
  const r = await fetch('/api/intents_stats').then(r=>r.json());
  const el = $('intStats');
  if(!r.stats.length){ el.innerHTML = '<p class="hint">Нет данных.</p>'; return; }
  el.innerHTML = '<div class="kpi">' + r.stats.map(s=>{const cls = s.percent>=50?'ok':(s.percent>=25?'warn':'err');
    return '<div>'+esc(s.nickname)+'<b>'+s.percent+'%</b><span class="hint"><span class="pill '+cls+'">'+s.matched+'/'+s.total+'</span></span></div>';
  }).join('') + '</div>';
}

async function loadCalendar(){
  const r = await fetch('/api/days?days=60').then(r=>r.json());
  const grid = $('calGrid'); grid.innerHTML = '';
  const byDay = {};
  r.days.forEach(d=>{ byDay[d.day] = d; });
  const today = new Date();
  const start = new Date(today); start.setDate(start.getDate() - 41);
  for (let i=0; i<42; i++) {
    const d = new Date(start); d.setDate(start.getDate() + i);
    const key = d.toISOString().slice(0,10);
    const dayData = byDay[key];
    const cell = document.createElement('div');
    cell.className = 'cal-day';
    if (dayData) {
      cell.innerHTML = '<div class="num">'+d.getDate()+'</div>'
        + '<div class="cnt">'+dayData.cnt+'</div>'
        + '<div class="rain">🌧'+(dayData.rains||0)+' 💸'+(dayData.tips||0)+'</div>';
      cell.onclick = () => loadDayDetail(key);
    } else {
      cell.innerHTML = '<div class="num">'+d.getDate()+'</div><div class="cnt">—</div>';
    }
    grid.appendChild(cell);
  }
  const docxInp = $('docxDate');
  if (docxInp && !docxInp.value) docxInp.value = new Date().toISOString().slice(0,10);
}
async function loadDayDetail(date){
  const el = $('calDetail');
  el.innerHTML = '<div class="vault-card">⏳ Загрузка...</div>';
  const r = await fetch('/api/day_detail?date='+date).then(r=>r.json());
  const recvHtml = r.top_receivers.map((p,i)=>'<tr><td>'+(i+1)+'</td><td><b>'+esc(p.nickname)+'</b></td><td>'+p.count+'</td><td style="color:#71d68a">'+fmtMoney(p.rub)+' ₽</td></tr>').join('') || '<tr><td colspan="4" class="hint">—</td></tr>';
  const sendHtml = r.top_senders.map((p,i)=>'<tr><td>'+(i+1)+'</td><td><b>'+esc(p.nickname)+'</b></td><td>'+p.count+'</td><td style="color:#71d68a">'+fmtMoney(p.rub)+' ₽</td></tr>').join('') || '<tr><td colspan="4" class="hint">—</td></tr>';
  el.innerHTML = '<div class="vault-card"><h3>📅 '+date+' — '+r.events_count+' событий</h3></div>'
    + '<div style="display:grid;grid-template-columns:1fr 1fr;gap:16px">'
    + '<div class="vault-card"><h3>🏆 Топ получателей</h3><table><thead><tr><th>#</th><th>Ник</th><th>Раз</th><th>₽</th></tr></thead><tbody>'+recvHtml+'</tbody></table></div>'
    + '<div class="vault-card"><h3>🌧 Топ раздающих</h3><table><thead><tr><th>#</th><th>Ник</th><th>Дождей</th><th>₽</th></tr></thead><tbody>'+sendHtml+'</tbody></table></div>'
    + '</div>';
}
function downloadDocx(){
  const d = $('docxDate').value;
  if (!d) { alert('Выбери дату'); return; }
  window.location.href = '/api/daily_report?date=' + encodeURIComponent(d);
}

let chart1=null, chart2=null;
async function loadChart(){
  const p = new URLSearchParams();
  if($('cType').value) p.set('type',$('cType').value);
  if($('cChat').value) p.set('chat',$('cChat').value);
  if($('cRange').value) p.set('last_hours',$('cRange').value);
  const r = await fetch('/api/chart?'+p).then(r=>r.json());
  const labels = [...Array(24).keys()].map(h=>String(h).padStart(2,'0')+':00');
  const counts = labels.map((_,h)=>r.buckets[h]?.count||0);
  const sums   = labels.map((_,h)=>r.buckets[h]?.sum||0);
  if(chart1) chart1.destroy(); if(chart2) chart2.destroy();
  chart1 = new Chart($('hourCount'), {type:'bar', data:{labels, datasets:[{label:'Кол-во', data:counts, backgroundColor:'#4a7ce0'}]}});
  chart2 = new Chart($('hourSum'), {type:'bar', data:{labels, datasets:[{label:'Сумма ₽', data:sums, backgroundColor:'#71d68a'}]}});
}

async function loadPromos(){
  const filter = $('promoFilter').value;
  const p = new URLSearchParams();
  if (filter === '0') p.set('only_unused','1');
  if (filter === '1') p.set('only_used','1');

  let r, stats;
  try {
    const [rr, ss] = await Promise.all([
      fetch('/api/promocodes?'+p).then(x=>x.json()),
      fetch('/api/promocodes/stats').then(x=>x.json()),
    ]);
    r = rr; stats = ss;
  } catch(e){
    $('promoList').innerHTML = '<p class="hint">❌ Ошибка: '+esc(e.message)+'</p>';
    return;
  }
  $('promoStats').textContent = `Всего: ${stats.total||0} | Исп: ${stats.used||0} | Осталось: ${stats.unused||0}`;

  const el = $('promoList');
  if (!r.promos || !r.promos.length){
    el.innerHTML = '<p class="hint">Пусто</p>';
    return;
  }
  el.innerHTML = r.promos.map(x => {
    const usedClass = x.used ? 'used' : '';
    const dateStr = x.discovered_at
      ? new Date(x.discovered_at*1000).toLocaleString('ru-RU')
      : (x.message_date || '—');
    const linkHtml = x.link
      ? `<a href="${esc(x.link)}" target="_blank" rel="noopener">🔗 Открыть в TG</a>`
      : '';
    const toggleLabel = x.used ? '↩️ Вернуть' : '✅ Использован';
    const vaultBtn = x.used_for_vault
      ? `<span class="pill ok">💎 В хранилище $${(x.vault_amount_usd||0).toFixed(2)}</span>`
      : `<button class="btn blue" onclick="promoToVault(${x.id})">💎 В хранилище</button>`;
    const statusPill = x.used
      ? '<span class="pill warn">Использован</span>'
      : '<span class="pill ok">Активен</span>';
    return `<div class="promo-card ${usedClass}">
      <div class="promo-head">
        <span class="promo-date">📅 ${esc(dateStr)}</span>
        ${statusPill}
      </div>
      <div class="promo-text">${esc(x.text)}</div>
      <div class="promo-actions">
        ${linkHtml}
        <button class="btn" onclick="promoToggle(${x.id})">${toggleLabel}</button>
        ${vaultBtn}
      </div>
    </div>`;
  }).join('');
}

async function promoToggle(id){
  try {
    const r = await fetch('/api/promo/toggle/'+id, {method:'POST'}).then(x=>x.json());
    if (r.ok) loadPromos();
  } catch(e){ alert('Ошибка: '+e.message); }
}

async function promoToVault(id){
  const amt = prompt('Сумма в USD:', '0');
  if (amt === null) return;
  const currency = (prompt('Валюта (USD/USDT):', 'USD') || 'USD').toUpperCase();
  try {
    const r = await fetch('/api/promo/to_vault/'+id, {
      method:'POST',
      headers:{'Content-Type':'application/json'},
      body: JSON.stringify({amount_usd: parseFloat(amt)||0, currency})
    }).then(x=>x.json());
    if (r.ok) loadPromos();
  } catch(e){ alert('Ошибка: '+e.message); }
}
async function testPromo(ev){
  const btn = (ev && ev.target) || document.querySelector('button[onclick*="testPromo"]');
  if (btn){ btn.disabled=true; btn.textContent='⏳'; }
  try {
    const r = await fetch('/api/promo/test',{method:'POST'}).then(x=>x.json());
    alert('Найдено: '+r.found+', новых: '+r.new);
    loadPromos();
  } catch(e){ alert('Ошибка: '+e.message); }
  finally { if (btn){ btn.disabled=false; btn.textContent='🔍 Проверить канал'; } }
}

async function loadActivityDashboard(){
  const win = $('actWindow').value || 60;
  $('actStatus').textContent = '⏳';
  try {
    const r = await fetch('/api/activity/dashboard?minutes=' + win).then(r=>r.json());
    const grid = $('actGrid');
    const wl = r.watchlist || [];
    if (!wl.length) { grid.innerHTML = '<p class="hint">Пусто</p>'; return; }
    const chats = ['ENGLISH','RUSSIAN','JAPANESE','POLISH','TURKISH','PORTUGUESE','SPANISH','FRENCH','CHINESE','KOREAN','VIETNAMESE','INDONESIAN'];
    let html = '<table><thead><tr><th>Игрок</th>';
    chats.forEach(c => html += '<th style="text-align:center;font-size:10px">' + c.substring(0,4) + '</th>');
    html += '<th>Всего</th></tr></thead><tbody>';
    wl.forEach(name => {
      const p = r.players[name] || { chats: {}, total_msgs: 0 };
      html += '<tr><td><b>' + esc(name) + '</b>' + (name === 'Merami' ? ' ⭐' : '') + '</td>';
      chats.forEach(c => {
        const info = p.chats[c];
        if (!info) { html += '<td style="text-align:center;color:#333">·</td>'; return; }
        const mins = info.minutes_ago;
        let color = '#666', symbol = '·';
        if (mins !== null && mins !== undefined) {
          if (mins <= 25) { color = '#4ade80'; symbol = mins + 'м'; }
          else if (mins <= 59) { color = '#fbbf24'; symbol = mins + 'м'; }
          else { color = '#f87171'; symbol = mins + 'м'; }
        }
        html += `<td style="text-align:center;color:${color}">${symbol}</td>`;
      });
      html += '<td>' + (p.total_msgs || 0) + '</td></tr>';
    });
    html += '</tbody></table>';
    grid.innerHTML = html;
    $('actStatus').textContent = 'обновлено: ' + new Date().toLocaleTimeString('ru-RU');
  } catch(e) { $('actStatus').textContent = '❌ ' + e.message; }
}

async function loadRainActivity(){
  const box = $('rainActivityBox');
  box.innerHTML = '<div class="vault-card">⏳ Загрузка...</div>';
  try {
    const r = await fetch('/api/rain_activity?hours=24&limit=50').then(r=>r.json());
    if (!r.rains.length) { box.innerHTML = '<p class="hint">Нет дождей за 24 часа</p>'; return; }
    box.innerHTML = '<h3 style="color:#8ab4f8">🌧 Расширенный отчёт (24ч, ' + r.count + ')</h3>'
      + r.rains.map(e => {
        const tstr = fmtTime(e.ts);
        const m = e.merami;
        let myStat = '';
        if (m.won) myStat = '<span class="pill ok">✅ ты выиграл #' + m.pos + '</span>';
        else if (m.minutes_ago !== null && m.minutes_ago !== undefined)
          myStat = m.was_active
            ? '<span class="pill warn">❌ не выиграл · активен ' + m.minutes_ago + 'м</span>'
            : '<span class="pill err">❌ не выиграл · ' + m.minutes_ago + 'м назад</span>';
        else myStat = '<span class="pill err">❌ не писал</span>';
        const top = (e.top20 || []).slice(0, 10).map(r => {
          const icon = r.wrote_in_25min ? '✅' : (r.in_chat_only ? '🟢' : '⏸');
          const ago = r.minutes_ago !== null ? r.minutes_ago + 'м' : '—';
          return '<span class="tag">' + icon + ' ' + esc(r.name) + ' ' + ago + '</span>';
        }).join(' ');
        return '<div class="vault-card" style="padding:10px">'
          + '<div style="display:flex;gap:10px;align-items:center;flex-wrap:wrap">'
          + '<span class="pill">' + esc(e.chat) + '</span>'
          + '<b>' + esc(e.sender) + '</b>'
          + '<span class="hint">' + tstr + '</span>'
          + '<span style="color:#71d68a">' + esc(e.amount_text) + ' ' + esc(e.crypto_symbol) + '</span>'
          + '<span class="hint">на ' + e.receivers_count + ' чел</span>'
          + myStat
          + '</div>'
          + '<div style="margin-top:6px;font-size:12px"><b>Топ активных:</b> ' + top + '</div>'
          + '</div>';
      }).join('');
  } catch(e) { box.innerHTML = '<p class="hint">❌ ' + e.message + '</p>'; }
}

async function loadModerators(){
  const r = await fetch('/api/moderators').then(r=>r.json());
  const body = document.getElementById('modBody'); body.innerHTML = '';
  if(!r.moderators || !r.moderators.length){
    body.innerHTML = '<tr><td colspan="5" class="hint">Пусто</td></tr>';
    document.getElementById('modStats').textContent = '';
    return;
  }
  const search = (document.getElementById('modSearch').value || '').toLowerCase();
  const filtered = r.moderators.filter(m => !search || (m.username||'').toLowerCase().includes(search));
  document.getElementById('modStats').textContent = 'Всего: ' + r.moderators.length;
  filtered.forEach(m => {
    const chats = Object.entries(m.chats||{}).map(([c,n]) => '<span class="tag">'+esc(c)+':'+n+'</span>').join(' ');
    const tr = document.createElement('tr');
    tr.innerHTML = '<td><b>'+esc(m.username)+'</b></td>'
      + '<td>'+(m.first_seen?fmtTime(m.first_seen):'—')+'</td>'
      + '<td>'+(m.last_seen?fmtTime(m.last_seen):'—')+'</td>'
      + '<td>'+chats+'</td>'
      + '<td><button class="btn" onclick="loadModMsgs(\''+esc(m.username)+'\')">📜</button></td>';
    body.appendChild(tr);
  });
}
async function loadModMsgs(username){
  const r = await fetch('/api/moderator_messages?username='+encodeURIComponent(username)+'&limit=100').then(r=>r.json());
  const el = document.getElementById('modMessages');
  if(!r.messages.length){ el.innerHTML = '<p class="hint">Нет сообщений</p>'; return; }
  el.innerHTML = '<h3 style="color:#8ab4f8">Последние сообщения '+esc(username)+'</h3>'
    + r.messages.map(m => '<div style="padding:6px 10px;border-left:3px solid #c58af9;margin:4px 0;font-size:12px">'
      + '<span class="hint">'+fmtTime(m.ts)+' · '+esc(m.chat)+'</span> · '+esc(m.text||'') + '</div>').join('');
}

function saveMeNick(){ const v = $('meNick').value.trim(); if(!v) return; localStorage.setItem('myNick', v); alert('✅ '+v); }
async function loadMe(){
  const un = $('meNick').value.trim(); if(!un) return;
  const s = $('meStatus'); if(s) s.textContent = '⏳...';
  try {
    const d = await fetch('/api/player/activity?username=' + encodeURIComponent(un) + '&windows=15,25,60,1440').then(r=>r.json());
    const wins = d.windows || [15,25,60,1440];
    const lbl = w => w === 1440 ? '24ч' : (w === 60 ? '1ч' : (w + 'м'));
    const totalRow = wins.map(w => {
      const v = d.total[w] || 0; const col = v > 0 ? '#4ade80' : '#666';
      return '<div style="background:#151821;border:1px solid #232a37;border-radius:8px;padding:10px 16px"><span class="hint" style="font-size:12px">За ' + lbl(w) + '</span><b style="display:block;font-size:22px;color:' + col + '">' + v + '</b></div>';
    }).join('');
    const rows = (d.chats || []).map(c => {
      const cells = wins.map(w => { const v = c.counts[w] || 0; return '<td style="text-align:center;color:' + (v>0?'#4ade80':'#555') + '">'+v+'</td>'; }).join('');
      const ll = c.minutes_ago === null ? '—' : (c.minutes_ago < 60 ? c.minutes_ago + 'м' : (c.minutes_ago < 1440 ? Math.floor(c.minutes_ago/60) + 'ч' : Math.floor(c.minutes_ago/1440) + 'д'));
      return '<tr><td><b>'+c.chat+'</b></td>'+cells+'<td>'+ll+'</td></tr>';
    }).join('');
    const head = '<tr><th>Чат</th>'+wins.map(w => '<th style="text-align:center">'+lbl(w)+'</th>').join('')+'<th>Последнее</th></tr>';
    $('meActivity').innerHTML = '<div style="display:flex;gap:10px;flex-wrap:wrap;margin-bottom:12px">'+totalRow+'</div>'
      + (rows ? '<table><thead>'+head+'</thead><tbody>'+rows+'</tbody></table>' : '<div class="hint">Нет чатов</div>');
    if(s) s.textContent = '✅';
    loadMeRains(); loadMeBets();
  } catch(e) { if(s) s.textContent = '❌ ' + e.message; }
}
async function loadMeRains(){
  const un = $('meNick').value.trim(); if(!un) return;
  const days = $('meRainDays') ? $('meRainDays').value : 30;
  const st = $('meRainsStatus'); if(st) st.textContent = '⏳';
  const el = $('meRains'); if(!el) return;
  const d = await fetch('/api/user/' + encodeURIComponent(un) + '/rains_received?days=' + days).then(r=>r.json());
  if(st) st.textContent = 'Найдено: ' + d.wins + ' | ' + fmtMoney(d.total_rub) + ' ₽ (~$' + fmtMoney(d.total_usd) + ')';
  if(!d.items || !d.items.length){ el.innerHTML = '<div class="hint">Нет</div>'; return; }
  el.innerHTML = '<table><thead><tr><th>Время</th><th>Чат</th><th>От</th><th>Сумма</th><th>Доля ₽</th><th>Доля $</th></tr></thead><tbody>'
    + d.items.map(i => '<tr><td>'+fmtTime(i.ts)+'</td><td>'+esc(i.chat)+'</td><td><b>'+esc(i.sender)+'</b></td><td>'+esc(i.amount_text)+' '+esc(i.crypto_symbol||'')+' <span class="hint">(÷'+i.receivers_count+')</span></td><td style="color:#71d68a;font-weight:600">'+fmtMoney(i.share_rub)+' ₽</td><td style="color:#8ab4f8">$'+fmtMoney(i.share_usd)+'</td></tr>').join('') + '</tbody></table>';
}
async function loadMeBets(){
  const un = $('meNick').value.trim(); if(!un) return;
  const el = $('meBets'); if(!el) return;
  const r = await fetch('/api/bets?sender=' + encodeURIComponent(un) + '&limit=20').then(r=>r.json());
  if(!r.bets || !r.bets.length){ el.innerHTML = '<div class="hint">Нет ставок</div>'; return; }
  el.innerHTML = '<table><thead><tr><th>Время</th><th>Чат</th><th>Игра</th><th>Множ.</th><th>Сумма</th><th>$</th></tr></thead><tbody>'
    + r.bets.map(b => '<tr><td>'+fmtTime(b.ts)+'</td><td>'+esc(b.chat)+'</td><td>'+esc(b.game)+'</td><td>'+esc(b.multiplier)+'</td><td>'+esc(b.amount_text)+' '+esc(b.crypto||'')+'</td><td style="color:#71d68a">$'+fmtMoney(b.amount_usd)+'</td></tr>').join('') + '</tbody></table>';
}
(function initMe(){ const saved = localStorage.getItem('myNick'); if(saved && $('meNick')) $('meNick').value = saved; })();

async function loadDb(){
  try {
    const [si, db] = await Promise.all([
      fetch('/api/server_info').then(r=>r.json()),
      fetch('/api/db_stats').then(r=>r.json()),
    ]);
    const started = new Date(si.started_at*1000);
    const upMin = Math.round(si.uptime_sec / 60);
    const ago = upMin < 60 ? upMin+'м' : Math.floor(upMin/60)+'ч '+(upMin%60)+'м';
    $('dbDeploy').innerHTML = 'Старт: <b>'+started.toLocaleString('ru-RU')+'</b> ('+ago+' назад)<br>commit: <b>'+esc(si.git_commit)+'</b> · branch: '+esc(si.git_branch)+'<br>Сервис: '+esc(si.service_name);
    const t = db.tables || {};
    let html = '<table><thead><tr><th>Таблица</th><th>Строк</th></tr></thead><tbody>';
    Object.keys(t).forEach(k => {
      html += '<tr><td><b>'+esc(k)+'</b></td><td>' + (t[k] === null ? '—' : t[k]) + '</td></tr>';
    });
    html += '</tbody></table>';
    html += '<p style="margin-top:10px">Дождей: <b>'+db.rains+'</b> · Чаевых: <b>'+db.tips+'</b></p>';
    $('dbStats').innerHTML = html;
  } catch(e){ $('dbStats').innerHTML = '❌ '+e.message; }
}

let _lastNgrok = '';
async function pollNgrok() {
  try {
    const r = await fetch('/api/ngrok_url').then(r=>r.json());
    if (r.url && r.url !== _lastNgrok) {
      _lastNgrok = r.url;
      document.getElementById('ngrokUrl').textContent = '🌐 ' + r.url;
    }
  } catch(e){}
}
function copyNgrok() {
  if (!_lastNgrok) return;
  if (navigator.clipboard) navigator.clipboard.writeText(_lastNgrok);
  alert('Скопировано: ' + _lastNgrok);
}
async function controlAction(action) {
  if (!confirm('Команда: ' + action + '?')) return;
  await fetch('/api/control/set', {method:'POST', headers:{'Content-Type':'application/json'}, body: JSON.stringify({action})});
  alert('✅ Отправлено: ' + action);
}
async function pollDeployInfo() {
  try {
    const r = await fetch('/api/server_info').then(r=>r.json());
    const el = document.getElementById('deployInfo');
    const started = new Date(r.started_at * 1000);
    const upMin = Math.round(r.uptime_sec / 60);
    const ago = upMin < 60 ? upMin + 'м' : Math.floor(upMin/60) + 'ч ' + (upMin % 60) + 'м';
    const dd = String(started.getDate()).padStart(2,'0');
    const mm = String(started.getMonth()+1).padStart(2,'0');
    const hh = String(started.getHours()).padStart(2,'0');
    const mi = String(started.getMinutes()).padStart(2,'0');
    el.textContent = '🚀 ' + dd + '.' + mm + ' ' + hh + ':' + mi + ' (' + ago + ')';
    el.title = 'Commit: ' + r.git_commit;
  } catch(e) {}
}

async function runAI(){
  const chat = $('aiChat').value; if(!chat){ alert('Выбери чат'); return; }
  const btn = event.target; btn.disabled=true; btn.textContent='⏳'; $('aiStatus').textContent = 'Отправляю...';
  const r = await fetch('/api/chat/request',{method:'POST',headers:{'Content-Type':'application/json'},body:JSON.stringify({chat})}).then(r=>r.json());
  const taskId = r.task_id; if(!taskId){ alert('Ошибка'); btn.disabled=false; return; }
  let tries = 0;
  const t = setInterval(async () => { tries++; const tt = await fetch('/api/chat/task/'+taskId).then(r=>r.json());
    if(tt.status === 'pending' || tt.status === 'reading'){ $('aiStatus').textContent = (tries*2)+'с'; }
    else if(tt.status === 'completed'){ clearInterval(t); $('aiStatus').textContent='✅'; btn.disabled=false; btn.textContent='🤖 Получить'; renderAIResult(tt.result, chat); }
    else if(tt.status === 'failed'){ clearInterval(t); $('aiStatus').textContent='❌'; btn.disabled=false; btn.textContent='🤖 Получить'; }
    if(tries > 60){ clearInterval(t); btn.disabled=false; btn.textContent='🤖 Получить'; }
  }, 2000);
}
function renderAIResult(result, chat){
  if(!result){ $('aiResult').innerHTML='<p class="hint">Нет</p>'; return; }
  $('aiTopic').innerHTML = result.topic ? '<div class="vault-card"><b>Тема:</b> '+esc(result.topic)+'</div>' : '';
  $('aiResult').innerHTML = '<div class="ai-grid"><div class="ai-col"><h3>💬 Сообщения</h3>'
    + (result.messages||[]).map(m=>'<div class="ai-msg" onclick="copyMsg(this,\''+esc(m).replace(/'/g,"&#39;")+'\')">'+esc(m)+'</div>').join('')
    + '</div><div class="ai-col"><h3>🙋 Фразы</h3>'
    + (result.phrases||[]).map(p=>'<div class="ai-phrase" onclick="copyMsg(this,\''+esc(p).replace(/'/g,"&#39;")+'\')">'+esc(p)+'</div>').join('')
    + '</div></div>';
}
function copyMsg(el, text){
  const doCopy = async () => { if(navigator.clipboard && window.isSecureContext){ try { await navigator.clipboard.writeText(text); return true; } catch(e){} }
    const ta = document.createElement('textarea'); ta.value = text; ta.style.position='fixed'; ta.style.opacity='0'; document.body.appendChild(ta);
    ta.focus(); ta.select(); try { document.execCommand('copy'); return true; } catch(e){ return false; } finally { document.body.removeChild(ta); } };
  doCopy().then(()=>{ el.classList.add('copied'); const o = el.textContent; el.textContent = '✅ '+o; setTimeout(()=>{ el.classList.remove('copied'); el.textContent = o; },1200); });
}
async function loadAILatest(){ const chat = $('aiChat').value || null; const p = new URLSearchParams(); if(chat) p.set('chat',chat); const r = await fetch('/api/chat/latest?'+p).then(r=>r.json()); if(!r.task){ $('aiResult').innerHTML='<p class="hint">Нет</p>'; return; } renderAIResult(r.task.result, r.task.chat); $('aiStatus').textContent = 'Последний: '+fmtTime(r.task.completed_at); }

async function loadPlayersTab(){
  const sum = await fetch('/api/players_by_chat').then(r=>r.json());
  const sEl = $('plChatsSummary');
  if(sum.chats && sum.chats.length){
    sEl.innerHTML = '<div class="kpi">' + sum.chats.map(c=>'<div>'+esc(c.chat)+'<b>'+c.total_players+'</b><span class="hint">🟢 '+c.online_players+' · '+fmtMoney(c.total_messages)+'</span></div>').join('') + '</div>';
  }
  const p = new URLSearchParams();
  if($('plChatFilter').value) p.set('chat',$('plChatFilter').value);
  if($('plUserFilter').value) p.set('username',$('plUserFilter').value);
  if($('plOnlineOnly').checked) p.set('online_only','1');
  p.set('limit','1000');
  const r = await fetch('/api/players?'+p).then(r=>r.json());
  const body = $('plPlayersBody'); body.innerHTML='';
  if(!r.players.length){ body.innerHTML = '<tr><td colspan="6" class="hint">Нет</td></tr>'; return; }
  const now = Date.now()/1000;
  r.players.forEach(x=>{
    const online = (now - (x.last_seen||0)) < 300;
    const tr = document.createElement('tr');
    tr.innerHTML = '<td data-v="'+esc(x.chat)+'"><b>'+esc(x.chat)+'</b></td><td data-v="'+esc(x.username)+'">'+esc(x.username)+'</td>'
      + '<td data-v="'+x.msg_count+'" style="color:#8ab4f8;font-weight:bold">'+x.msg_count+'</td>'
      + '<td data-v="'+(x.last_seen||0)+'">'+(x.last_seen?fmtTime(x.last_seen):'—')+'</td>'
      + '<td>'+(online?'<span class="pill ok">🟢</span>':'<span class="pill">off</span>')+'</td>'
      + '<td></td>';
    body.appendChild(tr);
  });
}
async function loadLanguages(){
  const r = await fetch('/api/players_by_language').then(r=>r.json());
  const f = ($('langFilter').value || '').toLowerCase();
  const players = (r.players||[]).filter(p => !f || (p.username||'').toLowerCase().includes(f));
  const el = $('langGrid');
  if(!players.length){ el.innerHTML = '<p class="hint">Нет</p>'; return; }
  el.innerHTML = players.slice(0,500).map(p=>'<div class="lang-card"><h4>'+esc(p.username)+'</h4>'
    + '<div class="hint">Сообщений: '+(p.total_msgs||0)+' · Чатов: '+(p.chat_count||0)+'</div>'
    + (p.chats||[]).map(c=>'<div class="player-row">'+esc(c)+'</div>').join('')+'</div>').join('');
}

loadChats();
refreshStatus();
pollNgrok();
pollDeployInfo();
setInterval(refreshStatus, 5000);
setInterval(pollNgrok, 10000);
setInterval(pollDeployInfo, 30000);
</script>
<div class="modal" id="imgModal" onclick="this.classList.remove('on')">
  <img id="imgModalImg" src="" alt="screenshot">
</div>
</body></html>
"""


@app.get("/monitor", response_class=HTMLResponse)
async def monitor_page():
    return MONITOR_HTML


@app.on_event("startup")
async def startup_event():
    _load_server_rates()
    try:
        init_db()
        print("[STARTUP] Таблицы БД готовы")
    except Exception as e:
        print(f"[STARTUP] Ошибка init_db: {e}")
    try:
        migrate_ts_once()
        print("[STARTUP] Миграция ts выполнена")
    except Exception as e:
        print(f"[STARTUP] migrate_ts_once: {e}")
    threading.Thread(target=rates_refresh_thread, daemon=True).start()
    threading.Thread(target=dedup_thread, daemon=True).start()
    threading.Thread(target=ticks_cleanup_thread, daemon=True).start()
    threading.Thread(target=promo_parser_thread, daemon=True).start()
    print(f"[PROMO] Поток @{PROMO_CHANNEL} запущен")
    print(f"[STARTUP] Сервер запущен, commit={SERVER_GIT_COMMIT}")


if __name__ == "__main__":
    import uvicorn
    uvicorn.run("main:app", host="0.0.0.0", port=int(os.environ.get("PORT", 8000)))
